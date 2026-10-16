from __future__ import annotations

import csv
import json
import math
import statistics
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


INPUT = Path("/Users/agatbe/Downloads/chinafintech/TIM - 11.xlsx")
OUT_DIR = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)/research_notes/China banks fintech Bloomberg/audit_tmp")


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def quarter_key(value):
    if not isinstance(value, datetime):
        return None
    return f"{value.year}Q{((value.month - 1) // 3) + 1}"


def pct_growth(current, prior):
    if not is_number(current) or not is_number(prior) or prior == 0:
        return None
    return (float(current) / float(prior) - 1.0) * 100.0


def median(values):
    clean = [float(v) for v in values if is_number(v)]
    return statistics.median(clean) if clean else None


wb_values = openpyxl.load_workbook(INPUT, data_only=True, read_only=False)
wb_formulas = openpyxl.load_workbook(INPUT, data_only=False, read_only=False)

date_sheet = wb_values["Bank 1"]
dates = []
for col in range(4, date_sheet.max_column + 1):
    value = date_sheet.cell(4, col).value
    if isinstance(value, datetime):
        dates.append((col, value))

latest_col_by_quarter = {}
dates_by_quarter = defaultdict(list)
for col, value in dates:
    key = quarter_key(value)
    dates_by_quarter[key].append((col, value))
    if key not in latest_col_by_quarter or value > latest_col_by_quarter[key][1]:
        latest_col_by_quarter[key] = (col, value)

quarter_columns = sorted(latest_col_by_quarter.items(), key=lambda item: item[1][1])

sheet_specs = {
    "Bank 1": {"sector": "bank", "block": 20},
    "Fin": {"sector": "fintech", "block": 15},
}

records = []
securities = []
metric_rows = defaultdict(dict)
for sheet_name, spec in sheet_specs.items():
    sheet = wb_values[sheet_name]
    row = 5
    while row <= sheet.max_row:
        ticker = sheet.cell(row, 1).value
        if not ticker:
            row += spec["block"]
            continue
        securities.append({"sheet": sheet_name, "sector": spec["sector"], "row": row, "ticker": ticker})
        for metric_row in range(row, min(row + spec["block"], sheet.max_row + 1)):
            metric = sheet.cell(metric_row, 3).value
            metric_rows[(sheet_name, ticker)][metric] = metric_row
            for quarter, (col, date_value) in quarter_columns:
                value = sheet.cell(metric_row, col).value
                records.append(
                    {
                        "sector": spec["sector"],
                        "sheet": sheet_name,
                        "ticker": ticker,
                        "metric": metric,
                        "quarter": quarter,
                        "observation_date": date_value.date().isoformat(),
                        "value": value,
                        "source_cell": f"{get_column_letter(col)}{metric_row}",
                        "is_numeric": is_number(value),
                        "missing_kind": None if is_number(value) else ("blank" if value in (None, "") else str(value)),
                    }
                )
        row += spec["block"]

with (OUT_DIR / "normalized_quarterly.csv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(records[0].keys()))
    writer.writeheader()
    writer.writerows(records)

values = {}
cell_sources = {}
for record in records:
    key = (record["sheet"], record["ticker"], record["metric"], record["quarter"])
    values[key] = record["value"]
    cell_sources[key] = record["source_cell"]

duplicate_tickers = {}
for sheet_name in sheet_specs:
    counts = Counter(item["ticker"] for item in securities if item["sheet"] == sheet_name)
    duplicate_tickers[sheet_name] = {ticker: count for ticker, count in counts.items() if count > 1}

coverage = []
for sheet_name, spec in sheet_specs.items():
    sheet_tickers = [item["ticker"] for item in securities if item["sheet"] == sheet_name]
    metrics = sorted({record["metric"] for record in records if record["sheet"] == sheet_name})
    for metric in metrics:
        relevant = [
            record
            for record in records
            if record["sheet"] == sheet_name and record["metric"] == metric
        ]
        numeric = [record for record in relevant if record["is_numeric"]]
        numeric_quarters = sorted({record["quarter"] for record in numeric})
        coverage.append(
            {
                "sheet": sheet_name,
                "metric": metric,
                "security_rows": len(sheet_tickers),
                "unique_securities": len(set(sheet_tickers)),
                "numeric_cells": len(numeric),
                "total_cells": len(relevant),
                "numeric_pct": 100 * len(numeric) / len(relevant) if relevant else None,
                "earliest_numeric_quarter": numeric_quarters[0] if numeric_quarters else None,
                "latest_numeric_quarter": numeric_quarters[-1] if numeric_quarters else None,
            }
        )


def company_snapshot(sheet_name, ticker, current_quarter, prior_quarter, metrics):
    result = {"ticker": ticker, "current_quarter": current_quarter, "prior_quarter": prior_quarter}
    for metric in metrics:
        current = values.get((sheet_name, ticker, metric, current_quarter))
        prior = values.get((sheet_name, ticker, metric, prior_quarter))
        result[metric] = current
        result[f"{metric}_prior"] = prior
        result[f"{metric}_growth_pct"] = pct_growth(current, prior)
        result[f"{metric}_cell"] = cell_sources.get((sheet_name, ticker, metric, current_quarter))
        result[f"{metric}_prior_cell"] = cell_sources.get((sheet_name, ticker, metric, prior_quarter))
    return result


bank_tickers = [item["ticker"] for item in securities if item["sheet"] == "Bank 1"]
bank_tickers_ex_cicc = [ticker for ticker in bank_tickers if ticker != "601995 CH Equity"]
fin_tickers_all = [item["ticker"] for item in securities if item["sheet"] == "Fin"]
fin_tickers_unique = list(dict.fromkeys(fin_tickers_all))

bank_metrics = [
    "BS_TOT_ASSET",
    "BS_GROSS_LOANS_COMPANY_BASIS",
    "BS_CUSTOMER_DEPOSITS",
    "TOTAL_EQUITY",
    "SALES_REV_TURN",
    "IS_OPER_INC",
    "IS_INT_INC",
    "ARD_NET_FEE_AND_COMMISSION_INC",
    "T12_NET_INT_MARGIN",
    "RETURN_ON_ASSET",
    "RETURN_COM_EQY",
    "NPLS_TO_TOTAL_LOANS",
    "IS_PROV_FOR_LOAN_LOSS",
    "BS_CET1_BUFFR_PCT",
    "PX_TO_BOOK_RATIO",
    "PE_RATIO",
    "BEST_PE_RATIO",
    "DIVIDEND_12_MONTH_YIELD",
]

fin_metrics = [
    "SALES_REV_TURN",
    "PG_REVENUE_VALUE",
    "EBIT",
    "EBITDA",
    "NET_INCOME",
    "BS_TOT_ASSET",
    "TOTAL_EQUITY",
    "RETURN_ON_ASSET",
    "RETURN_COM_EQY",
    "MONTHLY_ACTIVE_USERS",
    "PE_RATIO",
    "BEST_PE_RATIO",
    "PX_TO_BOOK_RATIO",
    "CURRENT_EV_TO_12M_SALES",
    "DIVIDEND_12_MONTH_YIELD",
]

snapshots = {
    "bank_2025": [company_snapshot("Bank 1", ticker, "2025Q4", "2024Q4", bank_metrics) for ticker in bank_tickers_ex_cicc],
    "bank_2026_h1": [company_snapshot("Bank 1", ticker, "2026Q2", "2025Q2", bank_metrics) for ticker in bank_tickers_ex_cicc],
    "fintech_2025": [company_snapshot("Fin", ticker, "2025Q4", "2024Q4", fin_metrics) for ticker in fin_tickers_unique],
    "fintech_2026_h1": [company_snapshot("Fin", ticker, "2026Q2", "2025Q2", fin_metrics) for ticker in fin_tickers_unique],
}


def summarize_snapshot(rows, level_metrics, growth_metrics):
    out = {}
    for metric in level_metrics:
        out[f"median_{metric}"] = median(row.get(metric) for row in rows)
        out[f"n_{metric}"] = len([row for row in rows if is_number(row.get(metric))])
    for metric in growth_metrics:
        field = f"{metric}_growth_pct"
        out[f"median_{field}"] = median(row.get(field) for row in rows)
        out[f"n_{field}"] = len([row for row in rows if is_number(row.get(field))])
    return out


summaries = {
    "bank_2025": summarize_snapshot(
        snapshots["bank_2025"],
        ["T12_NET_INT_MARGIN", "RETURN_ON_ASSET", "RETURN_COM_EQY", "NPLS_TO_TOTAL_LOANS", "PX_TO_BOOK_RATIO", "PE_RATIO", "DIVIDEND_12_MONTH_YIELD"],
        ["BS_TOT_ASSET", "BS_GROSS_LOANS_COMPANY_BASIS", "BS_CUSTOMER_DEPOSITS", "SALES_REV_TURN", "IS_OPER_INC"],
    ),
    "bank_2026_h1": summarize_snapshot(
        snapshots["bank_2026_h1"],
        ["T12_NET_INT_MARGIN", "RETURN_ON_ASSET", "RETURN_COM_EQY", "NPLS_TO_TOTAL_LOANS", "PX_TO_BOOK_RATIO", "PE_RATIO", "DIVIDEND_12_MONTH_YIELD"],
        ["BS_TOT_ASSET", "BS_GROSS_LOANS_COMPANY_BASIS", "BS_CUSTOMER_DEPOSITS", "SALES_REV_TURN", "IS_OPER_INC"],
    ),
    "fintech_2025": summarize_snapshot(
        snapshots["fintech_2025"],
        ["RETURN_ON_ASSET", "RETURN_COM_EQY", "PX_TO_BOOK_RATIO", "PE_RATIO", "BEST_PE_RATIO", "CURRENT_EV_TO_12M_SALES", "DIVIDEND_12_MONTH_YIELD"],
        ["BS_TOT_ASSET", "SALES_REV_TURN", "EBIT", "EBITDA", "NET_INCOME"],
    ),
    "fintech_2026_h1": summarize_snapshot(
        snapshots["fintech_2026_h1"],
        ["RETURN_ON_ASSET", "RETURN_COM_EQY", "PX_TO_BOOK_RATIO", "PE_RATIO", "BEST_PE_RATIO", "CURRENT_EV_TO_12M_SALES", "DIVIDEND_12_MONTH_YIELD"],
        ["BS_TOT_ASSET", "SALES_REV_TURN", "EBIT", "EBITDA", "NET_INCOME"],
    ),
}

formula_arrays = []
for sheet_name in ["Bank", "Fintech"]:
    sheet = wb_formulas[sheet_name]
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, openpyxl.worksheet.formula.ArrayFormula):
                formula_arrays.append(
                    {
                        "sheet": sheet_name,
                        "anchor": cell.coordinate,
                        "spill_ref": value.ref,
                        "formula": value.text,
                    }
                )

pair_differences = {}
for formula_sheet, value_sheet in [("Bank", "Bank 1"), ("Fintech", "Fin")]:
    left = wb_values[formula_sheet]
    right = wb_values[value_sheet]
    diffs = []
    for left_row, right_row in zip(
        left.iter_rows(min_row=4, min_col=3), right.iter_rows(min_row=4, min_col=3)
    ):
        for left_cell, right_cell in zip(left_row, right_row):
            if left_cell.value != right_cell.value:
                diffs.append(
                    {
                        "cell": left_cell.coordinate,
                        "formula_sheet_value": left_cell.value,
                        "static_sheet_value": right_cell.value,
                    }
                )
    pair_differences[f"{formula_sheet}_vs_{value_sheet}"] = diffs[:100]

jd_rows = [item["row"] for item in securities if item["sheet"] == "Fin" and item["ticker"] == "9618 HK Equity"]
jd_duplicate_mismatches = []
if len(jd_rows) == 2:
    first, second = jd_rows
    for offset in range(15):
        for col in range(1, wb_values["Fin"].max_column + 1):
            first_value = wb_values["Fin"].cell(first + offset, col).value
            second_value = wb_values["Fin"].cell(second + offset, col).value
            if col == 1 and offset == 0:
                continue
            if first_value != second_value:
                jd_duplicate_mismatches.append(
                    {
                        "first_cell": f"{get_column_letter(col)}{first + offset}",
                        "second_cell": f"{get_column_letter(col)}{second + offset}",
                        "first_value": first_value,
                        "second_value": second_value,
                    }
                )

date_summary = {
    "raw_date_columns": len(dates),
    "unique_calendar_quarters": len(quarter_columns),
    "earliest_date": min(value for _, value in dates).date().isoformat(),
    "latest_date": max(value for _, value in dates).date().isoformat(),
    "quarters_with_multiple_dates": {
        quarter: [value.date().isoformat() for _, value in items]
        for quarter, items in dates_by_quarter.items()
        if len(items) > 1
    },
}

error_counts = {}
for sheet_name in wb_values.sheetnames:
    sheet = wb_values[sheet_name]
    counts = Counter()
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("#"):
                counts[value] += 1
    error_counts[sheet_name] = dict(counts)

result = {
    "input": str(INPUT),
    "date_summary": date_summary,
    "securities": securities,
    "duplicate_tickers": duplicate_tickers,
    "coverage": coverage,
    "formula_arrays": formula_arrays,
    "pair_differences": pair_differences,
    "jd_duplicate_mismatches": jd_duplicate_mismatches[:100],
    "error_counts": error_counts,
    "snapshots": snapshots,
    "summaries": summaries,
}

(OUT_DIR / "analysis.json").write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
print(json.dumps({key: result[key] for key in ["date_summary", "duplicate_tickers", "error_counts", "summaries"]}, indent=2, default=str))
