from copy import deepcopy
from pathlib import Path

from docx import Document


SOURCE = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)/resume_edit/original.docx")
OUTPUT = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)/resume_edit/Resume - Agatbe Thymoty.docx")


REPLACEMENTS = {
    "LinkedIn | X/Twitter | Instagram | Portfolio": "linkedin.com/in/agatbe | agatbe.com",
    "PROFESSIONAL EXPERIENCE.": "PROFESSIONAL EXPERIENCE",
    "Investment research professional with experience across\u00a0wealth management, asset management, and equity research, spanning macroeconomics, fixed income, mutual funds, Indonesian equities, offshore equities, and digital assets. Currently at\u00a0BCA Wealth Management, supporting portfolio management through cross-asset research, daily and monthly market outlooks, and high-conviction investment product analysis. Strong foundation in\u00a0fundamental research, financial modeling, quantitative analysis, and portfolio construction, complemented by experience leading a 7-analyst research team and building a 4,600+ member financial research community..": (
        "Investment research and portfolio management professional with experience across wealth management, asset management, and equity research. "
        "Conduct cross-asset research covering macroeconomics, fixed income, mutual funds, Indonesian and offshore equities, and digital assets to support portfolio decisions and investment-product evaluation at BCA Wealth Management. "
        "Skilled in fundamental research, financial modeling, quantitative analysis, and portfolio construction, with additional experience leading a seven-analyst team and building a 4,600+ member financial research community."
    ),
    "Conducted\u00a0macroeconomic research\u00a0to assess global and domestic economic conditions, monetary policy, interest rates, inflation, and their implications for investment markets.": (
        "Analyze global and Indonesian macroeconomic conditions, including monetary policy, interest rates, and inflation, to assess cross-asset market implications."
    ),
    "Conducted research across\u00a0fixed income, mutual funds, offshore equities, and Indonesian equities\u00a0to support investment analysis and product evaluation.": (
        "Evaluate fixed income, mutual funds, Indonesian equities, and offshore equities to support investment-product screening and portfolio decisions."
    ),
    "Prepared\u00a0daily market outlooks\u00a0covering key overnight developments, building investment thesis, market movements, economic data, and potential implications for Indonesian investors. ": (
        "Produce daily market outlooks that synthesize overnight developments, market movements, economic data, and investment implications for Indonesian investors."
    ),
    "Assisted in\u00a0building AI software dashboard for portfolio and investment product analysis\u00a0by evaluating market conditions, asset-class performance, and relevant investment considerations for wealth management clients.": (
        "Contribute to an AI-enabled dashboard for portfolio and investment-product analysis by translating market conditions and asset-class performance into decision-relevant insights."
    ),
    "Designed 5 PowerPoint training modules translating Revalue Academy's REVALUE stock-selection framework into accessible financial lessons for live retail-investor classes.": (
        "Designed five PowerPoint training modules that translated Revalue Academy's REVALUE stock-selection framework into accessible lessons for live retail-investor classes."
    ),
    "Applied a visual-first, analogy-based approach to explain technical concepts such as net profit usage and dividend payout as intuitive visual narratives.": (
        "Used visual storytelling and analogies to explain net-income allocation and dividend policy to non-specialist audiences."
    ),
    "Converted technical valuation and capital-allocation concepts into clear, retail-facing content that improved comprehension without sacrificing accuracy.": (
        "Converted valuation and capital-allocation concepts into clear retail-investor content while preserving analytical rigor."
    ),
    "Led a student consulting team, sourced through SBM ITB Investment Club and ITB Student Consulting, to design an AI/NLP-based market intelligence framework for one of Indonesia's largest domestic asset managers.": (
        "Led a student consulting team from SBM ITB Investment Club and ITB Student Consulting to design an AI/NLP market-intelligence framework for a leading Indonesian asset manager."
    ),
    "Delivered a sentiment-aware methodology document and interactive dashboard prototype that automated financial news curation and surfaced early macroeconomic trend signals for the portfolio management team.": (
        "Developed a sentiment-analysis methodology and interactive dashboard prototype that automated financial-news curation and surfaced early macroeconomic signals for portfolio managers."
    ),
    "Handed off the proof-of-concept and methodology to the firm's investment team, providing an implementation roadmap for future production deployment.": (
        "Delivered the proof of concept, methodology, and implementation roadmap to the investment team for future production deployment."
    ),
    "Conducted bottom-up due diligence on 70+ DePIN (Decentralized Physical Infrastructure) projects, assessing tokenomics, ecosystem viability, and yield potential to support institutional allocation decisions.": (
        "Performed bottom-up due diligence on 70+ DePIN projects, evaluating tokenomics, ecosystem viability, and yield potential for institutional allocation decisions."
    ),
    "Authored monthly deal-flow reports covering 50+ crypto funding rounds, feeding directly into investment-thesis generation for VC and institutional fund clients.": (
        "Produced monthly deal-flow reports covering 50+ crypto funding rounds, informing investment-thesis development for VC and institutional-fund clients."
    ),
    "Mapped the Fantom network ecosystem through on-chain analysis, visualizing protocol interconnections and identifying emerging yield opportunities.": (
        "Mapped the Fantom ecosystem through on-chain analysis to visualize protocol relationships and identify emerging yield opportunities."
    ),
    "Degree conferred (Yudisium) August 2026; formal graduation ceremony October 2026.": (
        "Degree conferred (Yudisium) August 2026; formal graduation ceremony October 2026."
    ),
    "GPA: 3.45 / 4.00.": "GPA: 3.45/4.00.",
    "Run competitor analysis across executive education and business masterclass providers, benchmarking curriculum scope, delivery format, and positioning to inform how the program is differentiated.": (
        "Benchmark executive-education and business-masterclass providers across curriculum, delivery model, pricing, and positioning to inform program differentiation."
    ),
    "Develop pricing options for the CEO Masterclass, comparing tiered and bundled structures against competitor price points and the willingness to pay of the target segment.": (
        "Develop pricing options for the CEO Masterclass by evaluating tiered and bundled structures against competitor benchmarks and target-segment willingness to pay."
    ),
    "Support go-to-market planning by mapping target participant segments and outlining acquisition channels and launch sequencing with the program team.": (
        "Support go-to-market planning by defining target participant segments, acquisition channels, and launch sequencing with the program team."
    ),
    "Build and scale a Discord-based financial research community to 4,600+ members and 200+ paying subscribers, generating Rp50,000,000 in year-one program revenue through research subscriptions and paid investment classes averaging 500 attendees per session, ahead of full monetization.": (
        "Founded and scaled a Discord-based financial research community to 4,600+ members and 200+ paying subscribers, generating Rp50 million in first-year revenue through research subscriptions and investment classes averaging 500 attendees."
    ),
    "Grow distribution to 10,000+ followers across Instagram and X, sustaining consistent engagement and establishing brand recognition among Indonesian retail and student investors.": (
        "Built an audience of 10,000+ followers across Instagram and X, strengthening brand awareness among Indonesian retail and student investors."
    ),
    "Lead a 7-analyst research team through Notion-based workflows and weekly thesis reviews, publishing institutional-style research spanning Indonesian equities, macroeconomics, and digital assets.": (
        "Lead a seven-analyst team through structured Notion workflows and weekly thesis reviews, publishing institutional-style research across Indonesian equities, macroeconomics, and digital assets."
    ),
    "Close 20+ partnerships, including co-hosted webinars and cross-promotions with McKinsey YLI, Samuel Sekuritas, Syailendra Capital, and leading Indonesian universities": (
        "Secured 20+ partnerships and co-hosted webinars or cross-promotions with McKinsey YLI, Samuel Sekuritas, Syailendra Capital, and leading Indonesian universities."
    ),
    "LSP – Wakil Perantara Perdagangan Efek (WPPE): Issued April 2026 (Indonesian Broker-Dealer Representative License).": (
        "WPPE – Wakil Perantara Perdagangan Efek, LSP | Issued April 2026 | Indonesian Broker-Dealer Representative License"
    ),
}


RICH_REPLACEMENTS = {
    "Equity Research Competitions: Finalist & Top 4, IMPACT Maliki Equity Research Paper (2025, HOLD recommendation on PT Bank Central Asia Tbk – BBCA); Finalist & Best Speaker, UCMC Equity Research Competition.": (
        "Equity Research: ",
        "Top 4 Finalist, IMPACT Maliki Equity Research Paper 2025 (HOLD recommendation, PT Bank Central Asia Tbk/BBCA); Finalist and Best Speaker, UCMC Equity Research Competition.",
    ),
    "Entrepreneurial Grants: Selected as 1 of 3 recipients of the Ministry of Higher Education, Research, and Technology venture grant; 1 of 25 recipients in ITB Student Entrepreneurial Program.": (
        "Entrepreneurship: ",
        "Selected as 1 of 3 recipients of the Ministry of Higher Education, Research, and Technology venture grant and 1 of 25 recipients of the ITB Student Entrepreneurial Program.",
    ),
    'Guest Speaker: Delivered sessions on capital markets and investment topics at Lotus Andalan Sekuritas, Institut Teknologi Bandung, Universitas Indonesia, and UPN "Veteran".': (
        "Public Speaking: ",
        "Invited speaker on capital markets and investing at Lotus Andalan Sekuritas, Institut Teknologi Bandung, Universitas Indonesia, and UPN Veteran.",
    ),
    "Technical: Financial Modeling (DCF, Comparable Company Analysis, Dividend Discount Model), Equity Research, Portfolio Optimization (Markowitz), Monte Carlo Simulation, Quantitative Analysis, Financial Accounting Fundamentals.": (
        "Technical: ",
        "Financial Modeling (DCF, Comparable Company Analysis, Dividend Discount Model); Equity Research; Fundamental Analysis; Fixed-Income Analysis; Portfolio Construction and Optimization (Markowitz); Monte Carlo Simulation; Quantitative Analysis; Financial Statement Analysis; Macroeconomic Research.",
    ),
    "Tools: Microsoft Excel (advanced), PowerPoint, Python (pandas, yfinance, web scraping), Notion, Figma, Adobe Photoshop.": (
        "Tools: ",
        "Microsoft Excel (Advanced), PowerPoint, Python (pandas, yfinance, web scraping), Notion, Figma, Adobe Photoshop.",
    ),
    "Languages: Bahasa Indonesia (Native); English (Professional Working Proficiency).": (
        "Languages: ",
        "Bahasa Indonesia (Native); English (Professional Working Proficiency).",
    ),
}


def copy_run_format(source, target):
    if source is not None and source._r.rPr is not None:
        target._r.insert(0, deepcopy(source._r.rPr))


def replace_plain(paragraph, new_text):
    source_run = paragraph.runs[0] if paragraph.runs else None
    paragraph.clear()
    run = paragraph.add_run(new_text)
    copy_run_format(source_run, run)


def replace_rich(paragraph, label, body):
    normal_source = next((run for run in paragraph.runs if run.text and not run.bold), None)
    if normal_source is None and paragraph.runs:
        normal_source = paragraph.runs[-1]
    paragraph.clear()
    label_run = paragraph.add_run(label)
    copy_run_format(normal_source, label_run)
    label_run.bold = True
    body_run = paragraph.add_run(body)
    copy_run_format(normal_source, body_run)
    body_run.bold = False


doc = Document(SOURCE)
seen_plain = set()
seen_rich = set()

for paragraph in doc.paragraphs:
    original = paragraph.text
    if original in REPLACEMENTS:
        replace_plain(paragraph, REPLACEMENTS[original])
        seen_plain.add(original)
    elif original in RICH_REPLACEMENTS:
        replace_rich(paragraph, *RICH_REPLACEMENTS[original])
        seen_rich.add(original)

missing = (set(REPLACEMENTS) - seen_plain) | (set(RICH_REPLACEMENTS) - seen_rich)
if missing:
    raise RuntimeError(f"Expected text not found: {sorted(missing)!r}")

# Remove all consecutive empty spacer paragraphs before EDUCATION.
education_paragraph = next(p for p in doc.paragraphs if p.text == "EDUCATION")
previous = education_paragraph._p.getprevious()
while previous is not None and not "".join(previous.itertext()).strip():
    parent = previous.getparent()
    prior = previous.getprevious()
    parent.remove(previous)
    previous = prior

# Keep section headings with the following content to prevent orphaned headings.
section_headings = {
    "PROFESSIONAL SUMMARY",
    "PROFESSIONAL EXPERIENCE",
    "EDUCATION",
    "ORGANIZATIONAL EXPERIENCE",
    "CERTIFICATIONS",
    "ACHIEVEMENTS",
    "SKILLS & TOOLS",
}
for paragraph in doc.paragraphs:
    if paragraph.text in section_headings:
        paragraph.paragraph_format.keep_with_next = True

doc.save(OUTPUT)
print(OUTPUT)
