from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)")
SOURCE_ROOT = ROOT / "work/founders_live_final_v5"
SOURCE_LETTERS = SOURCE_ROOT / "letters_docx"
SOURCE_TOR = SOURCE_ROOT / "tor_docx/ToR Narasumber CEO Masterclass 2026.docx"
TOR_LOGO = SOURCE_ROOT / "tor_media/image2.png"
SIGNATURE = Path(
    "/var/folders/12/kf33qfk559sgsvh7429ht9bc0000gn/T/"
    "codex-clipboard-62c74574-3a0f-47b3-8786-390f0b93a477.png"
)

WORK = ROOT / "work/founders_live_final_v6"
UPDATED_LETTERS = WORK / "letters_docx"
COMBINED_DOCX = WORK / "combined_docx"
OUTPUT_DIR = ROOT / "output/docx/Undangan Narasumber CEO Masterclass - Revisi Final"
OUTPUT_ZIP = ROOT / "output/docx/Undangan Narasumber CEO Masterclass - Revisi Final DOCX.zip"


def replace_paragraph_text(paragraph, text: str, *, bold: bool | None = None) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        if bold is not None:
            paragraph.runs[0].bold = bold
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        run = paragraph.add_run(text)
        if bold is not None:
            run.bold = bold


def insert_signature(doc: Document, *, compact: bool = False) -> None:
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() != "Stanley Sebastian":
            continue

        if index > 0:
            doc.paragraphs[index - 1].paragraph_format.space_after = Pt(2)

        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)

        signature_run = paragraph.add_run()
        signature_run.add_picture(
            str(SIGNATURE), width=Inches(1.05 if compact else 1.55)
        )
        signature_run.add_break(WD_BREAK.LINE)
        name_run = paragraph.add_run("Stanley Sebastian Tjandra")
        name_run.bold = True
        name_run.font.name = "Calibri"
        name_run.font.size = Pt(11)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        return

    raise RuntimeError("Stanley Sebastian signature paragraph not found")


def personalize_otto_toto_sugiri(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Perihal"):
            replace_paragraph_text(
                paragraph, "Perihal\t:\tUndangan Narasumber Founder's Live", bold=True
            )
        elif text == "Pimpinan PT DCI Indonesia Tbk":
            replace_paragraph_text(paragraph, "Bapak Otto Toto Sugiri", bold=True)
        elif text == "u.p. [nama penerima]":
            replace_paragraph_text(paragraph, "Founder dan CEO DCI Indonesia")
        elif text.startswith("Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass"):
            replace_paragraph_text(
                paragraph,
                "Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass pada November "
                "2026, kami bermaksud mengundang Bapak sebagai narasumber pada sesi yang akan "
                "dilaksanakan pada:",
            )
        elif text.startswith("CEO Masterclass adalah bagian dari Kamar Dagang"):
            replace_paragraph_text(
                paragraph,
                "CEO Masterclass adalah bagian dari Kamar Dagang dan Industri (KADIN) Indonesia "
                "dan merupakan program pembelajaran bisnis yang materinya dibawakan oleh para "
                "pemimpin usaha nasional, antara lain Anindya Bakrie, Shinta Kamdani, Sandiaga "
                "Uno, dan Pandu Sjahrir. Founder's Live terdiri atas delapan sesi daring yang "
                "mempertemukan pelaku usaha dengan anggota CEO Masterclass untuk membahas "
                "pengalaman dalam menjalankan bisnis. Dalam kurikulum CEO Masterclass, materi "
                "Modul 6 Operation & Legal dibawakan oleh Bapak bersama Anne Patricia Sutanto. "
                "Kami berharap Bapak berkenan membagikan pengalaman membangun dan mengelola DCI "
                "Indonesia kepada peserta.",
            )
        elif text.startswith("Sesi berlangsung selama 90 menit"):
            replace_paragraph_text(
                paragraph,
                "Setiap sesi berlangsung selama 90 menit dan dipandu oleh moderator dari CEO "
                "Masterclass. Narasumber membagikan satu pengalaman atau studi kasus selama "
                "kurang lebih 20 menit, kemudian berdiskusi dengan peserta. Materi presentasi "
                "tidak wajib, dan informasi perusahaan yang bersifat rahasia tidak perlu "
                "disampaikan. Koordinasi sebelum acara dilakukan bersama Bapak melalui chat "
                "WhatsApp.",
            )
        elif text.startswith("Untuk koordinasi lebih lanjut"):
            replace_paragraph_text(
                paragraph,
                "Untuk koordinasi lebih lanjut, Bapak dapat menghubungi Official CEO Masterclass "
                "melalui WhatsApp di +6285178906100.",
            )
        elif text.startswith("Demikian permohonan ini kami sampaikan"):
            replace_paragraph_text(
                paragraph,
                "Demikian surat undangan ini kami sampaikan. Atas perhatian dan kesediaan Bapak, "
                "kami ucapkan terima kasih.",
            )


def update_letters() -> list[Path]:
    UPDATED_LETTERS.mkdir(parents=True, exist_ok=True)
    for old_file in UPDATED_LETTERS.glob("*.docx"):
        old_file.unlink()

    outputs: list[Path] = []
    for source in sorted(SOURCE_LETTERS.glob("*.docx")):
        output_name = source.name
        if "Pimpinan DCI Indonesia" in source.name:
            output_name = "Undangan Narasumber CEO Masterclass Bapak Otto Toto Sugiri.docx"

        target = UPDATED_LETTERS / output_name
        shutil.copy2(source, target)
        doc = Document(target)
        if "Pimpinan DCI Indonesia" in source.name:
            personalize_otto_toto_sugiri(doc)
        compact_signature = any(
            name in source.name for name in ("Pandu Sjahrir", "Pimpinan Saratoga", "Pimpinan Sinar Mas")
        )
        insert_signature(doc, compact=compact_signature)
        doc.save(target)
        outputs.append(target)

    if len(outputs) != 9:
        raise RuntimeError(f"Expected 9 updated letters, found {len(outputs)}")
    return outputs


def remove_body_children_from(body, start_element) -> None:
    removing = False
    for child in list(body):
        if child is start_element:
            removing = True
        if removing and child.tag != qn("w:sectPr"):
            body.remove(child)


def body_content(document: Document):
    for child in document._body._body:
        if child.tag != qn("w:sectPr"):
            yield child


def remove_page_breaks(element) -> None:
    for page_break in list(element.xpath('.//w:br[@w:type="page"]')):
        parent = page_break.getparent()
        if parent is not None:
            parent.remove(page_break)


def clear_header_or_footer(part) -> None:
    element = part._element
    for child in list(element):
        element.remove(child)
    element.append(OxmlElement("w:p"))


def add_bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), "D9D9D9")
    borders.append(bottom)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction_text, separate, placeholder, end])


def configure_tor_section(section, source_section) -> None:
    for attribute in (
        "page_width",
        "page_height",
        "top_margin",
        "bottom_margin",
        "left_margin",
        "right_margin",
        "header_distance",
        "footer_distance",
    ):
        setattr(section, attribute, getattr(source_section, attribute))

    section.header.is_linked_to_previous = False
    clear_header_or_footer(section.header)
    header_paragraph = section.header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_paragraph.add_run().add_picture(str(TOR_LOGO), width=Inches(2.55))
    add_bottom_border(header_paragraph)

    section.footer.is_linked_to_previous = False
    clear_header_or_footer(section.footer)
    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.paragraph_format.space_after = Pt(0)
    left_run = footer_paragraph.add_run("CEO Masterclass · ceomasterclass.id")
    left_run.font.name = "Aptos"
    left_run.font.size = Pt(8)


def copy_letter_section_format(target_section, source_section) -> None:
    for attribute in (
        "page_width",
        "page_height",
        "top_margin",
        "bottom_margin",
        "left_margin",
        "right_margin",
        "header_distance",
        "footer_distance",
    ):
        setattr(target_section, attribute, getattr(source_section, attribute))

    for reference in list(target_section._sectPr.findall(qn("w:headerReference"))):
        target_section._sectPr.remove(reference)
    for reference in source_section._sectPr.findall(qn("w:headerReference")):
        target_section._sectPr.insert(0, deepcopy(reference))

    target_section.footer.is_linked_to_previous = False
    clear_header_or_footer(target_section.footer)

    for page_numbering in list(target_section._sectPr.findall(qn("w:pgNumType"))):
        target_section._sectPr.remove(page_numbering)


def combine_letter_with_tor(letter_path: Path, tor: Document, output_path: Path) -> None:
    source_letter = Document(letter_path)
    target = Document(letter_path)

    appendix_source = next(
        paragraph for paragraph in source_letter.paragraphs if paragraph.text.strip() == "LAMPIRAN 1"
    )
    appendix_elements = []
    collecting = False
    for child in source_letter._body._body:
        if child is appendix_source._p:
            collecting = True
        if collecting and child.tag != qn("w:sectPr"):
            appendix_elements.append(deepcopy(child))
    if not appendix_elements:
        raise RuntimeError(f"Appendix content not found in {letter_path.name}")
    remove_page_breaks(appendix_elements[0])

    appendix_target = next(
        paragraph for paragraph in target.paragraphs if paragraph.text.strip() == "LAMPIRAN 1"
    )
    remove_body_children_from(target._body._body, appendix_target._p)

    letter_section_source = source_letter.sections[0]
    tor_section_source = tor.sections[0]

    tor_section = target.add_section(WD_SECTION.NEW_PAGE)
    configure_tor_section(tor_section, tor_section_source)
    for child in body_content(tor):
        target._body._body.insert(-1, deepcopy(child))

    appendix_section = target.add_section(WD_SECTION.NEW_PAGE)
    copy_letter_section_format(appendix_section, letter_section_source)
    for child in appendix_elements:
        target._body._body.insert(-1, child)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    target.save(output_path)


def build_combined_docx(updated_letters: list[Path]) -> list[Path]:
    COMBINED_DOCX.mkdir(parents=True, exist_ok=True)
    for old_file in COMBINED_DOCX.glob("*.docx"):
        old_file.unlink()

    tor = Document(SOURCE_TOR)
    outputs: list[Path] = []
    for letter in updated_letters:
        output = COMBINED_DOCX / letter.name
        combine_letter_with_tor(letter, tor, output)
        outputs.append(output)
    return outputs


def package_outputs(combined_files: list[Path]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for old_file in OUTPUT_DIR.glob("*.docx"):
        old_file.unlink()
    if OUTPUT_ZIP.exists():
        OUTPUT_ZIP.unlink()

    output_files: list[Path] = []
    for source in combined_files:
        target = OUTPUT_DIR / source.name
        shutil.copy2(source, target)
        output_files.append(target)

    standalone_tor = OUTPUT_DIR / SOURCE_TOR.name
    shutil.copy2(SOURCE_TOR, standalone_tor)
    output_files.append(standalone_tor)

    with zipfile.ZipFile(OUTPUT_ZIP, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for file in sorted(output_files):
            archive.write(file, arcname=file.name)

    print(f"Created {len(output_files)} DOCX files in {OUTPUT_DIR}")
    print(f"Created ZIP: {OUTPUT_ZIP}")


def main() -> None:
    if not SIGNATURE.exists():
        raise FileNotFoundError(SIGNATURE)
    updated_letters = update_letters()
    combined_files = build_combined_docx(updated_letters)
    package_outputs(combined_files)


if __name__ == "__main__":
    main()
