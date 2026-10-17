from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)")
SOURCE_LETTERS = ROOT / "work/founders_live_final_v3/letters_docx"
SOURCE_TOR = (
    ROOT
    / "work/founders_live_final_v4/tor_docx"
    / "CMC - ToR Founders Live 2026 - Umum.docx"
)
OUT_ROOT = ROOT / "work/founders_live_final_v5"
LETTER_OUT = OUT_ROOT / "letters_docx"
TOR_OUT = OUT_ROOT / "tor_docx" / "ToR Narasumber CEO Masterclass 2026.docx"

GREEN = "285B4A"
LIGHT_GREEN = "E8F0ED"
LIGHT_GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)


LETTER_CONFIG = {
    "Richard Theodore": {
        "output": "Undangan Narasumber CEO Masterclass Bapak Richard Theodore.docx",
        "honorific": "Bapak",
        "points": (
            "Kapan bisnis perlu bertahan, kapan perlu mengubah arah, dan pelajaran dari keputusan "
            "yang tidak berhasil."
        ),
    },
    "Yukka Harlanda": {
        "output": "Undangan Narasumber CEO Masterclass Bapak Yukka Harlanda.docx",
        "honorific": "Bapak",
        "points": "Menentukan masalah utama pelanggan dan memilih solusi yang perlu diuji lebih dulu.",
    },
    "Merry Riana": {
        "output": "Undangan Narasumber CEO Masterclass Ibu Merry Riana.docx",
        "honorific": "Ibu",
        "points": "Kapan mulai membangun tim, siapa yang direkrut lebih dulu, dan cara mendelegasikan pekerjaan.",
    },
    "dr. Richard Lee": {
        "output": "Undangan Narasumber CEO Masterclass Bapak dr. Richard Lee.docx",
        "honorific": "Bapak",
        "points": "Kapan mulai membangun tim, siapa yang direkrut lebih dulu, dan cara mendelegasikan pekerjaan.",
    },
    "Ryan Maurice Tallulah": {
        "output": "Undangan Narasumber CEO Masterclass Bapak Ryan Maurice Tallulah.docx",
        "honorific": "Bapak",
        "points": "Menguji masalah pelanggan, membuat produk awal, dan menyusun nilai jual yang jelas.",
    },
    "CEO Office Sinar Mas": {
        "output": "Undangan Narasumber CEO Masterclass Pimpinan Sinar Mas.docx",
        "honorific": None,
        "points": "Memilih kanal penjualan, menetapkan harga, dan membuat proses penjualan yang dapat diulang.",
    },
    "CEO Office DCI Indonesia": {
        "output": "Undangan Narasumber CEO Masterclass Pimpinan DCI Indonesia.docx",
        "honorific": None,
        "points": "Menata SOP, legalitas, kontrak, dan proses kerja utama agar operasional lebih rapi.",
    },
    "CEO Office Saratoga": {
        "output": "Undangan Narasumber CEO Masterclass Pimpinan Saratoga.docx",
        "honorific": None,
        "points": "Memahami arus kas, modal kerja, dan penyebab bisnis untung tetapi kekurangan kas.",
    },
    "CEO Office Pandu Sjahrir": {
        "output": "Undangan Narasumber CEO Masterclass Bapak Pandu Sjahrir.docx",
        "honorific": "Bapak",
        "points": "Menilai kesiapan bisnis untuk scale-up dan memilih sumber pendanaan yang sesuai.",
    },
}


SCHEDULE = [
    (
        "Selasa, 17 November",
        "Modul 1  Entrepreneurship Mentality",
        "Kapan bisnis perlu bertahan, kapan perlu mengubah arah, dan pelajaran dari keputusan yang tidak berhasil.",
    ),
    (
        "Kamis, 19 November",
        "Modul 2  Problem-Solving",
        "Menentukan masalah utama pelanggan dan memilih solusi yang perlu diuji lebih dulu.",
    ),
    (
        "Selasa, 24 November",
        "Modul 3  Team and Networking Formation",
        "Kapan mulai membangun tim, siapa yang direkrut lebih dulu, dan cara mendelegasikan pekerjaan.",
    ),
    (
        "Kamis, 26 November",
        "Modul 4  Minimum Sellable Product",
        "Menguji masalah pelanggan, membuat produk awal, dan menyusun nilai jual yang jelas.",
    ),
    (
        "Selasa, 1 Desember",
        "Modul 5  Go-To-Market",
        "Memilih kanal penjualan, menetapkan harga, dan membuat proses penjualan yang dapat diulang.",
    ),
    (
        "Kamis, 3 Desember",
        "Modul 6  Operation and Legal",
        "Menata SOP, legalitas, kontrak, dan proses kerja utama agar operasional lebih rapi.",
    ),
    (
        "Selasa, 8 Desember",
        "Modul 7  Finance and Tax",
        "Memahami arus kas, modal kerja, dan penyebab bisnis untung tetapi kekurangan kas.",
    ),
    (
        "Kamis, 10 Desember",
        "Modul 8  Scale Up and Fundraising",
        "Menilai kesiapan bisnis untuk scale-up dan memilih sumber pendanaan yang sesuai.",
    ),
]


def set_paragraph_text_keep_format(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_clone_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_xml = deepcopy(paragraph._p)
    paragraph._p.addnext(new_xml)
    new_paragraph = Paragraph(new_xml, paragraph._parent)
    set_paragraph_text_keep_format(new_paragraph, text)
    return new_paragraph


def remove_highlights(doc: Document) -> None:
    for element in doc.element.xpath(".//w:highlight"):
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def find_letter_config(filename: str) -> dict:
    for marker, config in LETTER_CONFIG.items():
        if marker in filename:
            return config
    raise RuntimeError(f"No letter configuration for {filename}")


def update_letter(source: Path, target: Path, config: dict) -> None:
    shutil.copy2(source, target)
    doc = Document(target)
    remove_highlights(doc)

    honorific = config["honorific"]
    if honorific:
        for paragraph in doc.paragraphs:
            if "Bapak/Ibu" in paragraph.text:
                set_paragraph_text_keep_format(
                    paragraph, paragraph.text.replace("Bapak/Ibu", honorific)
                )

    topic_paragraph = next(
        (p for p in doc.paragraphs if p.text.strip().startswith("Topik")), None
    )
    if topic_paragraph is None:
        raise RuntimeError(f"Topic paragraph not found in {source.name}")
    points_paragraph = insert_clone_after(
        topic_paragraph, f"Poin bahasan\t:\t{config['points']}"
    )
    points_paragraph.paragraph_format.space_after = Pt(5)

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Setiap sesi berlangsung selama 90 menit"):
            salutation = honorific or "perwakilan yang ditunjuk"
            text = (
                "Setiap sesi berlangsung selama 90 menit dan dipandu oleh moderator dari CEO "
                "Masterclass. Narasumber membagikan satu pengalaman atau studi kasus selama "
                "kurang lebih 20 menit, kemudian berdiskusi dengan peserta. Materi presentasi "
                "tidak wajib, dan informasi perusahaan yang bersifat rahasia tidak perlu "
                f"disampaikan. Koordinasi sebelum acara dilakukan bersama {salutation} melalui "
                "chat WhatsApp."
            )
            set_paragraph_text_keep_format(paragraph, text)
        elif paragraph.text.startswith("Sesi berlangsung selama 90 menit"):
            text = (
                "Sesi berlangsung selama 90 menit dan dipandu oleh moderator dari CEO "
                "Masterclass. Perwakilan yang ditunjuk membagikan satu pengalaman atau studi "
                "kasus selama kurang lebih 20 menit, kemudian berdiskusi dengan peserta. Materi "
                "presentasi tidak wajib, dan informasi perusahaan yang bersifat rahasia tidak "
                "perlu disampaikan. Koordinasi sebelum acara dilakukan melalui chat WhatsApp."
            )
            set_paragraph_text_keep_format(paragraph, text)

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Untuk koordinasi lebih lanjut"):
            subject = honorific or "perwakilan"
            set_paragraph_text_keep_format(
                paragraph,
                f"Untuk koordinasi lebih lanjut, {subject} dapat menghubungi Official CEO "
                "Masterclass melalui WhatsApp di +6285178906100.",
            )
            break
    else:
        raise RuntimeError(f"Contact paragraph not found in {source.name}")

    doc.save(target)


def build_letters() -> list[Path]:
    LETTER_OUT.mkdir(parents=True, exist_ok=True)
    for old_file in LETTER_OUT.glob("*.docx"):
        old_file.unlink()

    outputs = []
    for source in sorted(SOURCE_LETTERS.glob("*.docx")):
        config = find_letter_config(source.name)
        target = LETTER_OUT / config["output"]
        update_letter(source, target, config)
        outputs.append(target)
    if len(outputs) != 9:
        raise RuntimeError(f"Expected 9 letters, found {len(outputs)}")
    return outputs


def set_run_font(run, size: float, bold: bool = False, color=BLACK) -> None:
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_text_paragraph(
    doc: Document,
    text: str,
    *,
    size: float = 9.5,
    bold: bool = False,
    color=BLACK,
    before: float = 0,
    after: float = 4,
    keep_with_next: bool = False,
) -> Paragraph:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.keep_with_next = keep_with_next
    set_run_font(paragraph.add_run(text), size, bold, color)
    return paragraph


def add_section_heading(doc: Document, text: str) -> Paragraph:
    return add_text_paragraph(
        doc,
        text,
        size=11.2,
        bold=True,
        color=RGBColor(40, 91, 74),
        before=8,
        after=3,
        keep_with_next=True,
    )


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_GRAY, size="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def write_cell(cell, text: str, *, size=9, bold=False, color=BLACK) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.02
    set_run_font(paragraph.add_run(text), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_info_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.42)
    table.columns[1].width = Inches(5.45)
    for row, (label, value) in zip(table.rows, rows):
        set_row_cant_split(row)
        write_cell(row.cells[0], label, bold=True)
        write_cell(row.cells[1], value)
        for cell in row.cells:
            set_cell_borders(cell, color="E1E1E1", size="3")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_grid_table(
    doc: Document,
    headers: tuple[str, ...],
    rows: list[tuple[str, ...]],
    widths: tuple[float, ...],
    *,
    body_size: float = 8.8,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for index, (header, width) in enumerate(zip(headers, widths)):
        table.columns[index].width = Inches(width)
        cell = table.rows[0].cells[index]
        write_cell(cell, header, size=8.8, bold=True, color=RGBColor(40, 91, 74))
        set_cell_shading(cell, LIGHT_GREEN)
        set_cell_borders(cell)

    for values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            write_cell(cell, value, size=body_size, bold=(index == 0 and len(headers) == 2))
            set_cell_borders(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.20)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.04
    set_run_font(paragraph.add_run("• "), 9.2)
    set_run_font(paragraph.add_run(text), 9.2)


def clear_document_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def build_tor() -> Path:
    TOR_OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE_TOR, TOR_OUT)
    doc = Document(TOR_OUT)
    clear_document_body(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.86)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.76)
    section.right_margin = Inches(0.76)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK

    add_text_paragraph(
        doc,
        "KERANGKA ACUAN KEGIATAN (ToR)",
        size=8.6,
        bold=True,
        color=RGBColor(40, 91, 74),
        after=1,
        keep_with_next=True,
    )
    add_text_paragraph(
        doc,
        "Founder's Live CEO Masterclass 2026",
        size=17.2,
        bold=True,
        color=RGBColor(40, 91, 74),
        after=8,
        keep_with_next=True,
    )

    add_section_heading(doc, "1. Kegiatan")
    add_text_paragraph(
        doc,
        "Founder's Live adalah rangkaian sesi daring CEO Masterclass yang mempertemukan member "
        "dengan founder dan pemimpin bisnis. Setiap narasumber membagikan satu pengalaman yang "
        "berkaitan dengan modul yang sedang dipelajari, kemudian mendiskusikannya bersama peserta.",
    )
    add_text_paragraph(
        doc,
        "CEO Masterclass merupakan bagian dari Kamar Dagang dan Industri (KADIN) Indonesia. "
        "Kegiatan ini terdiri atas delapan sesi yang dilaksanakan setiap Selasa dan Kamis mulai "
        "17 November sampai 10 Desember 2026.",
        after=5,
    )

    add_section_heading(doc, "2. Pelaksanaan")
    add_info_table(
        doc,
        [
            ("Tanggal", "17 November-10 Desember 2026, setiap Selasa dan Kamis"),
            ("Waktu", "19.00-20.30 WIB"),
            ("Tempat", "Daring melalui Zoom"),
            ("Peserta", "Member CEO Masterclass, dengan kapasitas hingga 500 peserta"),
        ],
    )

    add_section_heading(doc, "3. Susunan sesi")
    add_grid_table(
        doc,
        ("Durasi", "Bagian", "Pelaksanaan"),
        [
            ("5 menit", "Pembukaan", "Moderator membuka sesi dan memperkenalkan narasumber."),
            ("20 menit", "Cerita narasumber", "Satu pengalaman atau studi kasus yang sesuai dengan topik."),
            ("55 menit", "Diskusi", "Tanya jawab dan pembahasan kasus peserta bersama moderator."),
            ("10 menit", "Penutup", "Rangkuman singkat, kuis, dan dokumentasi."),
        ],
        (0.92, 1.42, 4.45),
        body_size=8.9,
    )

    add_section_heading(doc, "4. Peran narasumber")
    add_bullet(doc, "Briefing dan koordinasi dilakukan melalui chat WhatsApp sebelum hari acara.")
    add_bullet(doc, "Masuk ke ruang Zoom 15 menit sebelum sesi untuk pemeriksaan audio dan video.")
    add_bullet(doc, "Membagikan satu pengalaman atau studi kasus selama sekitar 20 menit. Slide tidak wajib.")
    add_bullet(doc, "Menjawab pertanyaan yang dipilih moderator. Informasi rahasia perusahaan tidak perlu disampaikan.")

    doc.add_page_break()

    add_section_heading(doc, "5. Jadwal dan poin bahasan")
    add_text_paragraph(
        doc,
        "Poin bahasan berikut menjadi arah diskusi dan dapat disesuaikan dengan pengalaman narasumber.",
        size=9.1,
        after=4,
    )
    add_grid_table(
        doc,
        ("Tanggal", "Modul", "Bahasan singkat"),
        SCHEDULE,
        (1.26, 2.16, 3.37),
        body_size=8.2,
    )

    add_section_heading(doc, "6. Rekaman dan honorarium")
    add_text_paragraph(
        doc,
        "Sesi direkam untuk materi pembelajaran member dan publikasi program. Penggunaan nama, "
        "foto, dan cuplikan sesi mengikuti persetujuan tertulis dari narasumber.",
        size=9.1,
        after=3,
    )
    add_text_paragraph(
        doc,
        "Honorarium diberikan sesuai kesepakatan tertulis. Nilai, pajak, dan waktu pembayaran "
        "disampaikan secara terpisah sebelum acara.",
        size=9.1,
        after=3,
    )

    add_section_heading(doc, "7. Kontak")
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    set_run_font(paragraph.add_run("Official CEO Masterclass: "), 9.4, bold=True)
    set_run_font(paragraph.add_run("+6285178906100"), 9.4)

    remove_highlights(doc)
    doc.save(TOR_OUT)
    return TOR_OUT


def main() -> None:
    letters = build_letters()
    tor = build_tor()
    print(f"Created {len(letters)} revised invitation DOCX files")
    print(tor)


if __name__ == "__main__":
    main()
