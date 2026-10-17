from pathlib import Path
import shutil

from docx import Document


ROOT = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)")
SOURCE_DIR = ROOT / "output/docx/Undangan Narasumber CEO Masterclass - Revisi Final"
WORK_DIR = ROOT / "work/founders_live_recapital_pandu_team"


def replace_text(paragraph, text: str, *, bold: bool | None = None) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        if bold is not None:
            paragraph.runs[0].bold = bold
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        run = paragraph.add_run(text)
        if bold is not None:
            run.bold = bold


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def build_recapital() -> Path:
    source = SOURCE_DIR / "Undangan Narasumber CEO Masterclass Pimpinan Saratoga.docx"
    output = WORK_DIR / "Undangan Narasumber CEO Masterclass RE Capital.docx"
    shutil.copy2(source, output)
    document = Document(output)

    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("Nomor"):
            replace_text(paragraph, "Nomor\t:\tCMC/044/UND/X/2026")
        elif text == "Pimpinan PT Saratoga Investama Sedaya Tbk":
            replace_text(paragraph, "Pimpinan RE Capital", bold=True)
        elif text == "u.p. [nama penerima]":
            remove_paragraph(paragraph)
        elif text.startswith("Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass"):
            replace_text(
                paragraph,
                "Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass pada November "
                "2026, kami bermaksud memohon kesediaan RE Capital untuk menugaskan satu orang "
                "perwakilan sebagai narasumber pada sesi yang akan dilaksanakan pada:",
            )
        elif text.startswith("CEO Masterclass adalah bagian dari Kamar Dagang"):
            replace_text(
                paragraph,
                "CEO Masterclass adalah bagian dari Kamar Dagang dan Industri (KADIN) Indonesia "
                "dan merupakan program pembelajaran bisnis yang materinya dibawakan oleh para "
                "pemimpin usaha nasional, antara lain Anindya Bakrie, Shinta Kamdani, Sandiaga "
                "Uno, dan Pandu Sjahrir. Founder's Live terdiri atas delapan sesi daring yang "
                "mempertemukan pelaku usaha dengan anggota CEO Masterclass untuk membahas "
                "pengalaman dalam menjalankan bisnis. Dalam kurikulum CEO Masterclass, materi "
                "Modul 7 Finance & Tax dibawakan oleh Bapak Sandiaga Uno. Oleh karena itu, kami "
                "berharap RE Capital berkenan menugaskan perwakilan, misalnya Chief of Staff, "
                "Head of CEO Office, atau anggota tim strategi dan keuangan, untuk mengisi sesi "
                "tersebut.",
            )
        elif text.startswith("Untuk koordinasi lebih lanjut"):
            replace_text(
                paragraph,
                "Untuk koordinasi lebih lanjut, tim RE Capital dapat menghubungi Official CEO "
                "Masterclass melalui WhatsApp di +6285178906100.",
            )
        elif text.startswith("Demikian permohonan ini kami sampaikan"):
            replace_text(
                paragraph,
                "Demikian permohonan ini kami sampaikan. Atas perhatian dan dukungan RE Capital, "
                "kami ucapkan terima kasih.",
            )
        elif "Materi Modul 7 Finance & Tax dibawakan oleh Sandiaga Uno (Saratoga)." in text:
            replace_text(
                paragraph,
                text.replace("Sandiaga Uno (Saratoga)", "Sandiaga Uno (RE Capital)"),
            )

    document.save(output)
    return output


def build_pandu_team() -> Path:
    source = SOURCE_DIR / "Undangan Narasumber CEO Masterclass Bapak Pandu Sjahrir.docx"
    output = WORK_DIR / "Undangan Narasumber CEO Masterclass Tim Bapak Pandu Sjahrir.docx"
    shutil.copy2(source, output)
    document = Document(output)

    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("Nomor"):
            replace_text(paragraph, "Nomor\t:\tCMC/045/UND/X/2026")
        elif text == "Bapak Pandu Sjahrir":
            replace_text(paragraph, "Tim Bapak Pandu Sjahrir", bold=True)
        elif text == "u.p. [nama penerima]":
            remove_paragraph(paragraph)
        elif text.startswith("Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass"):
            replace_text(
                paragraph,
                "Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass pada November "
                "2026, kami bermaksud memohon kesediaan Tim Bapak Pandu Sjahrir untuk menunjuk "
                "satu orang perwakilan sebagai narasumber pada sesi yang akan dilaksanakan pada:",
            )
        elif text.startswith("CEO Masterclass adalah bagian dari Kamar Dagang"):
            replace_text(
                paragraph,
                "CEO Masterclass adalah bagian dari Kamar Dagang dan Industri (KADIN) Indonesia "
                "dan merupakan program pembelajaran bisnis yang materinya dibawakan oleh para "
                "pemimpin usaha nasional, antara lain Anindya Bakrie, Shinta Kamdani, Sandiaga "
                "Uno, dan Pandu Sjahrir. Founder's Live terdiri atas delapan sesi daring yang "
                "mempertemukan pelaku usaha dengan anggota CEO Masterclass untuk membahas "
                "pengalaman dalam menjalankan bisnis. Dalam kurikulum CEO Masterclass, materi "
                "Modul 8 Scale Up & Fundraising dibawakan oleh Bapak Pandu Sjahrir bersama "
                "Sandiaga Uno. Oleh karena itu, kami berharap sesi untuk modul ini dapat diisi "
                "oleh perwakilan Tim Bapak Pandu Sjahrir, misalnya Chief of Staff, Head of CEO "
                "Office, atau anggota tim investasi.",
            )
        elif text.startswith("Untuk koordinasi lebih lanjut"):
            replace_text(
                paragraph,
                "Untuk koordinasi lebih lanjut, Tim Bapak Pandu Sjahrir dapat menghubungi "
                "Official CEO Masterclass melalui WhatsApp di +6285178906100.",
            )
        elif text.startswith("Demikian permohonan ini kami sampaikan"):
            replace_text(
                paragraph,
                "Demikian permohonan ini kami sampaikan. Atas perhatian dan dukungan Tim Bapak "
                "Pandu Sjahrir, kami ucapkan terima kasih.",
            )

    document.save(output)
    return output


def main() -> None:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    outputs = [build_recapital(), build_pandu_team()]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
