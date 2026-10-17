from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)")
SOURCE = (
    ROOT
    / "output/docx/Undangan Narasumber CEO Masterclass - Revisi Final"
    / "Undangan Narasumber CEO Masterclass Bapak Otto Toto Sugiri.docx"
)
WORK_DIR = ROOT / "work/founders_live_dci_revision"
OUTPUT_DOCX = WORK_DIR / "Undangan Narasumber CEO Masterclass PT DCI Indonesia Tbk.docx"
SIGNATURE = Path(
    "/var/folders/12/kf33qfk559sgsvh7429ht9bc0000gn/T/"
    "codex-clipboard-62c74574-3a0f-47b3-8786-390f0b93a477.png"
)


def replace_text(paragraph, text: str, *, bold: bool | None = None) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        if bold is not None:
            paragraph.runs[0].bold = bold
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        run = paragraph.add_run(text)
        if bold is not None:
            run.bold = bold


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def compact_signature(document: Document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.strip() != "Stanley Sebastian Tjandra":
            continue
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)
        signature_run = paragraph.add_run()
        signature_run.add_picture(str(SIGNATURE), width=Inches(1.05))
        signature_run.add_break(WD_BREAK.LINE)
        name_run = paragraph.add_run("Stanley Sebastian Tjandra")
        name_run.bold = True
        name_run.font.name = "Calibri"
        name_run.font.size = Pt(11)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        return
    raise RuntimeError("Signature paragraph not found")


def main() -> None:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT_DOCX)
    document = Document(OUTPUT_DOCX)

    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text == "Bapak Otto Toto Sugiri":
            replace_text(paragraph, "Pimpinan PT DCI Indonesia Tbk", bold=True)
        elif text == "Founder dan CEO DCI Indonesia":
            remove_paragraph(paragraph)
        elif text.startswith("Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass"):
            replace_text(
                paragraph,
                "Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass pada November "
                "2026, kami bermaksud memohon kesediaan PT DCI Indonesia Tbk untuk menugaskan "
                "satu orang perwakilan sebagai narasumber pada sesi yang akan dilaksanakan pada:",
            )
        elif text.startswith("CEO Masterclass adalah bagian dari Kamar Dagang"):
            replace_text(
                paragraph,
                "CEO Masterclass adalah bagian dari Kamar Dagang dan Industri (KADIN) Indonesia "
                "dan merupakan program pembelajaran bisnis yang materinya dibawakan oleh para "
                "pemimpin usaha nasional, antara lain Anindya Bakrie, Shinta Kamdani, Sandiaga "
                "Uno, dan Pandu Sjahrir. Founder's Live terdiri atas delapan sesi daring yang "
                "mempertemukan pelaku usaha dengan anggota CEO Masterclass untuk membahas "
                "pengalaman dalam menjalankan bisnis. Dalam kurikulum CEO Masterclass, materi "
                "Modul 6 Operation & Legal dibawakan oleh Bapak Otto Toto Sugiri bersama Anne "
                "Patricia Sutanto. Kami berharap DCI Indonesia berkenan menugaskan perwakilan "
                "yang dapat membagikan pengalaman membangun dan mengelola DCI Indonesia kepada "
                "peserta. Apabila Bapak Otto Toto Sugiri berkenan hadir langsung sebagai "
                "narasumber, kami dengan senang hati akan menyambut kehadiran beliau.",
            )
        elif text.startswith("Setiap sesi berlangsung selama 90 menit"):
            replace_text(
                paragraph,
                "Setiap sesi berlangsung selama 90 menit dan dipandu oleh moderator dari CEO "
                "Masterclass. Narasumber yang ditunjuk membagikan satu pengalaman atau studi "
                "kasus selama kurang lebih 20 menit, kemudian berdiskusi dengan peserta. Materi "
                "presentasi tidak wajib, dan informasi perusahaan yang bersifat rahasia tidak "
                "perlu disampaikan. Koordinasi sebelum acara dilakukan bersama perwakilan yang "
                "ditunjuk melalui chat WhatsApp.",
            )
        elif text.startswith("Untuk koordinasi lebih lanjut"):
            replace_text(
                paragraph,
                "Untuk koordinasi lebih lanjut, tim DCI Indonesia dapat menghubungi Official CEO "
                "Masterclass melalui WhatsApp di +6285178906100.",
            )
        elif text.startswith("Demikian surat undangan ini kami sampaikan"):
            replace_text(
                paragraph,
                "Demikian surat undangan ini kami sampaikan. Atas perhatian dan dukungan PT DCI "
                "Indonesia Tbk, kami ucapkan terima kasih.",
            )

    compact_signature(document)
    document.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
