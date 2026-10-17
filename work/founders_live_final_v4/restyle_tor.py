from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor


ROOT = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)")
SOURCE = ROOT / "work/founders_live_final_v3/tor_docx/CMC - ToR Founders Live 2026 - Umum.docx"
OUTPUT = ROOT / "work/founders_live_final_v4/tor_docx/CMC - ToR Founders Live 2026 - Umum.docx"


def remove_cell_shading(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for shading in list(tc_pr.findall(qn("w:shd"))):
        tc_pr.remove(shading)


def set_cell_borders_none(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")


def set_table_borders_none(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")


def restyle_table(table) -> None:
    set_table_borders_none(table)
    for row in table.rows:
        for cell in row.cells:
            remove_cell_shading(cell)
            set_cell_borders_none(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = 0
                paragraph.paragraph_format.space_after = 0
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    for table in doc.tables:
        restyle_table(table)
    doc.save(OUTPUT)
    print(f"Restyled {len(doc.tables)} ToR tables: {OUTPUT}")


if __name__ == "__main__":
    main()
