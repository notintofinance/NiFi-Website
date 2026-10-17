from pathlib import Path
import shutil

from docx import Document


ROOT = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)")
SOURCE = (
    ROOT
    / "work/founders_live_recapital_pandu_team"
    / "Undangan Narasumber CEO Masterclass RE Capital.docx"
)
OUTPUT = (
    ROOT
    / "output/docx"
    / "Undangan Narasumber CEO Masterclass PT Samala Serasi Utama.docx"
)


def replace_text(paragraph, text: str, *, bold: bool | None = None) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        if bold is not None:
            paragraph.runs[0].bold = bold
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        run = paragraph.add_run(text)
        if bold is not None:
            run.bold = bold


def build() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text.startswith("Nomor"):
            replace_text(paragraph, "Nomor\t:\tCMC/046/UND/X/2026")
        elif text == "Pimpinan RE Capital":
            replace_text(paragraph, "Pimpinan PT Samala Serasi Utama", bold=True)
        elif text.startswith("Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass"):
            replace_text(
                paragraph,
                "Sehubungan dengan penyelenggaraan Founder's Live CEO Masterclass pada November "
                "2026, kami bermaksud memohon kesediaan PT Samala Serasi Utama untuk menghadirkan "
                "Bapak Sandiaga Uno atau menugaskan satu orang perwakilan dari tim beliau sebagai "
                "narasumber pada sesi yang akan dilaksanakan pada:",
            )
        elif text.startswith("CEO Masterclass adalah bagian dari Kamar Dagang"):
            replace_text(
                paragraph,
                "CEO Masterclass adalah bagian dari Kamar Dagang dan Industri (KADIN) Indonesia "
                "dan merupakan program pembelajaran bisnis yang materinya dibawakan oleh para "
                "pemimpin usaha nasional, antara lain Anindya Bakrie, Shinta Kamdani, Sandiaga "
                "Uno, dan Pandu Sjahrir. Founder's Live terdiri atas delapan sesi daring yang "
                "mempertemukan pelaku usaha dengan anggota CEO Masterclass untuk membahas "
                "pengalaman dalam menjalankan bisnis. Dalam kurikulum CEO Masterclass, materi "
                "Modul 7 Finance & Tax dibawakan oleh Bapak Sandiaga Uno. Oleh karena itu, kami "
                "berharap Bapak Sandiaga Uno atau perwakilan dari PT Samala Serasi Utama, misalnya "
                "Chief of Staff, Head of CEO Office, atau anggota tim strategi dan keuangan, "
                "berkenan mengisi sesi tersebut.",
            )
        elif text.startswith("Untuk koordinasi lebih lanjut"):
            replace_text(
                paragraph,
                "Untuk koordinasi lebih lanjut, Tim Samala Serasi Utama dapat menghubungi Official "
                "CEO Masterclass melalui WhatsApp di +6285178906100.",
            )
        elif text.startswith("Demikian permohonan ini kami sampaikan"):
            replace_text(
                paragraph,
                "Demikian permohonan ini kami sampaikan. Atas perhatian dan dukungan PT Samala "
                "Serasi Utama, kami ucapkan terima kasih.",
            )
        elif "Materi Modul 7 Finance & Tax dibawakan oleh Sandiaga Uno (RE Capital)." in text:
            replace_text(
                paragraph,
                text.replace(
                    "Sandiaga Uno (RE Capital)",
                    "Sandiaga Uno (Samala Serasi Utama)",
                ),
            )

    document.core_properties.title = (
        "Undangan Narasumber CEO Masterclass PT Samala Serasi Utama"
    )
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
