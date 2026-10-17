from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    "/Users/agatbe/Claude Project/Claude Vibecode (NiFi)/output/docx/"
    "Broadcast WhatsApp Undangan Pembicara.docx"
)


MESSAGES = [
    (
        "Broadcast untuk Semua Speaker",
        [
            "Halo Bapak/Ibu, perkenalkan kami dari Tim CEO Masterclass, powered by KADIN Indonesia, sebuah platform edukasi bisnis bagi para founder dan pemimpin usaha.",
            "Mewakili Founder kami, Bapak Stanley Sebastian Tjandra, kami ingin mengundang Bapak/Ibu sebagai pembicara dalam sesi kelas online Founder's Live CEO Masterclass.",
            "Surat undangan dan detail kegiatan kami lampirkan pada pesan ini. Besar harapan kami Bapak/Ibu dapat berkenan hadir dan berbagi pengalaman kepada para peserta CEO Masterclass.",
            "Terima kasih atas perhatian dan kesediaan Bapak/Ibu.",
            "Salam,\nTim CEO Masterclass",
        ],
    ),
]


def set_cell_margins(cell, top=110, start=150, bottom=110, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Arial", size=11.0, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_borders(paragraph_or_style):
    element = paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_message(doc, number, recipient, paragraphs, show_number=True):
    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(13)
    heading.paragraph_format.space_after = Pt(6)
    heading.add_run(f"{number}. {recipient}" if show_number else recipient)

    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.9)
    cell = table.cell(0, 0)
    set_cell_margins(cell)
    set_cell_border(cell)

    first = cell.paragraphs[0]
    for idx, text in enumerate(paragraphs):
        p = first if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(7 if idx < len(paragraphs) - 1 else 0)
        p.paragraph_format.line_spacing = 1.08
        for part_idx, part in enumerate(text.split("\n")):
            if part_idx:
                p.add_run().add_break()
            run = p.add_run(part)
            set_run_font(run)

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.line_spacing = 1


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(21)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_borders(title_style)

    heading_style = styles["Heading 1"]
    heading_style.font.name = "Arial"
    heading_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading_style.font.size = Pt(13)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(6)
    title.add_run("Broadcast WhatsApp Undangan Pembicara")
    remove_paragraph_borders(title)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run(
        "Pesan broadcast untuk mengundang seluruh pembicara ke sesi kelas online Founder's Live CEO Masterclass."
    )
    set_run_font(run, size=10.5, color="404040")

    for index, (recipient, paragraphs) in enumerate(MESSAGES, start=1):
        add_message(doc, index, recipient, paragraphs, show_number=len(MESSAGES) > 1)

    doc.core_properties.title = "Broadcast WhatsApp Undangan Pembicara"
    doc.core_properties.subject = "Undangan pembicara Founder's Live CEO Masterclass"
    doc.core_properties.author = "Tim CEO Masterclass"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
