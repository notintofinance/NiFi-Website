from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    "/Users/agatbe/Claude Project/Claude Vibecode (NiFi)/output/docx/"
    "Chat WhatsApp Individual Semua Speaker CEO Masterclass.docx"
)

PHONE = "+62 851-7890-6100"


SPEAKERS = [
    {
        "heading": "Bapak Richard Theodore",
        "type": "direct",
        "greeting": "Halo Bapak Richard Theodore,",
        "pronoun": "Bapak",
        "short_name": "Bapak Richard",
        "date": "Selasa, 17 November 2026",
        "topic": 'Modul 1 Entrepreneurship Mentality - "Kapan Harus Bertahan, Kapan Harus Berubah Arah?"',
    },
    {
        "heading": "Bapak Yukka Harlanda",
        "type": "direct",
        "greeting": "Halo Bapak Yukka Harlanda,",
        "pronoun": "Bapak",
        "short_name": "Bapak Yukka",
        "date": "Kamis, 19 November 2026",
        "topic": 'Modul 2 Problem-Solving - "Masalah Mana yang Harus Diselesaikan Dulu?"',
    },
    {
        "heading": "Ibu Merry Riana",
        "type": "direct",
        "greeting": "Halo Ibu Merry Riana,",
        "pronoun": "Ibu",
        "short_name": "Ibu Merry",
        "date": "Selasa, 24 November 2026",
        "topic": 'Modul 3 Team & Networking Formation - "Kapan Harus Mulai Bangun Team?"',
    },
    {
        "heading": "Bapak dr. Richard Lee",
        "type": "direct",
        "greeting": "Halo Bapak dr. Richard Lee,",
        "pronoun": "Bapak",
        "short_name": "Bapak Richard",
        "date": "Selasa, 24 November 2026",
        "topic": 'Modul 3 Team & Networking Formation - "Kapan Harus Mulai Bangun Team?"',
    },
    {
        "heading": "Bapak Ryan Maurice Tallulah",
        "type": "direct",
        "greeting": "Halo Bapak Ryan Maurice Tallulah,",
        "pronoun": "Bapak",
        "short_name": "Bapak Ryan",
        "date": "Kamis, 26 November 2026",
        "topic": 'Modul 4 Minimum Sellable Product - "Apa yang Bikin Produk Benar-Benar Layak Dijual?"',
    },
    {
        "heading": "Tim CEO Office Sinar Mas",
        "type": "office",
        "greeting": "Halo Bapak/Ibu Tim CEO Office Sinar Mas,",
        "invitee": "pimpinan Sinar Mas atau perwakilan dari CEO Office yang ditunjuk",
        "date": "Selasa, 1 Desember 2026",
        "topic": 'Modul 5 Go-To Market - "Gimana Menemukan Cara Jual yang Efektif?"',
    },
    {
        "heading": "Tim DCI Indonesia",
        "type": "office",
        "greeting": "Halo Bapak/Ibu Tim DCI Indonesia,",
        "invitee": "Bapak Otto Toto Sugiri atau perwakilan dari DCI Indonesia yang ditunjuk",
        "date": "Kamis, 3 Desember 2026",
        "topic": 'Modul 6 Operation & Legal - "Kapan Bisnis Harus Mulai Dibikin Lebih Rapi?"',
    },
    {
        "heading": "Tim RE Capital",
        "type": "office",
        "greeting": "Halo Bapak/Ibu Tim RE Capital,",
        "invitee": "Bapak Sandiaga Uno atau perwakilan dari RE Capital yang ditunjuk",
        "date": "Selasa, 8 Desember 2026",
        "topic": 'Modul 7 Finance & Tax - "Bisnis Untung, Tapi Kenapa Cash Liquid Nggak Ada?"',
    },
    {
        "heading": "Tim CEO Office Saratoga",
        "type": "office",
        "greeting": "Halo Bapak/Ibu Tim CEO Office Saratoga,",
        "invitee": "Bapak Sandiaga Uno atau perwakilan dari Saratoga yang ditunjuk",
        "date": "Selasa, 8 Desember 2026",
        "topic": 'Modul 7 Finance & Tax - "Bisnis Untung, Tapi Kenapa Cash Liquid Nggak Ada?"',
    },
    {
        "heading": "Tim Bapak Pandu Sjahrir",
        "type": "office",
        "greeting": "Halo Bapak/Ibu Tim Bapak Pandu Sjahrir,",
        "invitee": "Bapak Pandu Sjahrir atau perwakilan dari tim Bapak Pandu yang ditunjuk",
        "date": "Kamis, 10 Desember 2026",
        "topic": 'Modul 8 Scale Up & Fundraising - "Kapan Bisnis Sudah Siap untuk Scale dan Cari Modal?"',
    },
]


def set_run_font(run, size=11, bold=False, italic=False, color="000000"):
    run.font.name = "Arial"
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), "Arial")
    r_pr.rFonts.set(qn("w:hAnsi"), "Arial")
    r_pr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_borders(paragraph_or_style):
    element = paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_body_paragraph(doc, text, *, space_after=5, keep_with_next=True):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.06
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if "🙏" in text:
        before, after = text.split("🙏", 1)
        set_run_font(paragraph.add_run(before))
        emoji_run = paragraph.add_run("🙏")
        emoji_run.font.name = "Apple Color Emoji"
        emoji_r_pr = emoji_run._element.get_or_add_rPr()
        emoji_r_pr.rFonts.set(qn("w:ascii"), "Apple Color Emoji")
        emoji_r_pr.rFonts.set(qn("w:hAnsi"), "Apple Color Emoji")
        emoji_r_pr.rFonts.set(qn("w:eastAsia"), "Apple Color Emoji")
        emoji_run.font.size = Pt(10.5)
        if after:
            set_run_font(paragraph.add_run(after))
    else:
        set_run_font(paragraph.add_run(text))
    return paragraph


def add_schedule(doc, date, topic):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True

    labels_and_values = [
        ("Hari/tanggal: ", date),
        ("Waktu: ", "19.00-20.30 WIB"),
        ("Topik: ", topic),
    ]
    for index, (label, value) in enumerate(labels_and_values):
        if index:
            paragraph.add_run().add_break()
        label_run = paragraph.add_run(label)
        set_run_font(label_run, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run)


def add_message(doc, number, speaker):
    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.keep_with_next = True
    heading.add_run(f"{number}. {speaker['heading']}")

    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(0)
    label.paragraph_format.space_after = Pt(6)
    label.paragraph_format.keep_with_next = True
    label_text = (
        "Pesan langsung kepada narasumber"
        if speaker["type"] == "direct"
        else "Pesan melalui tim atau CEO Office"
    )
    set_run_font(label.add_run(label_text), size=9.5, italic=True, color="595959")

    add_body_paragraph(
        doc,
        speaker["greeting"]
        + " perkenalkan kami dari tim CEO Masterclass, platform edukasi bisnis yang powered by KADIN Indonesia.",
    )

    if speaker["type"] == "direct":
        add_body_paragraph(
            doc,
            "Mewakili Founder kami, Bapak Stanley Sebastian Tjandra, kami ingin mengundang "
            + speaker["pronoun"]
            + " sebagai pembicara dalam sesi kelas online Founder's Live CEO Masterclass pada:",
        )
    else:
        add_body_paragraph(
            doc,
            "Mewakili Founder kami, Bapak Stanley Sebastian Tjandra, kami ingin mengundang "
            + speaker["invitee"]
            + " untuk menjadi pembicara dalam sesi kelas online Founder's Live CEO Masterclass.",
        )

    add_schedule(doc, speaker["date"], speaker["topic"])

    if speaker["type"] == "direct":
        add_body_paragraph(
            doc,
            "Kami sangat berharap "
            + speaker["pronoun"]
            + " berkenan membagikan pengalaman dan wawasan kepada "
            "peserta kami. Surat undangan dan ToR kami lampirkan untuk informasi lebih lengkap.",
        )
        add_body_paragraph(
            doc,
            "Apakah " + speaker["pronoun"] + " berkenan menjadi pembicara pada sesi tersebut?",
        )
        add_body_paragraph(
            doc,
            "Terima kasih, " + speaker["short_name"] + ". 🙏",
        )
    else:
        add_body_paragraph(
            doc,
            "Mohon bantuan tim untuk menyampaikan dan mengoordinasikan undangan ini. Surat "
            "undangan dan ToR kami lampirkan sebagai informasi lebih lengkap.",
        )
        add_body_paragraph(
            doc,
            "Terima kasih atas perhatian dan bantuannya. Kami menantikan kabar baik dari Bapak/Ibu. 🙏",
        )

    signature = add_body_paragraph(doc, "Salam,\nTim CEO Masterclass\n" + PHONE, space_after=0, keep_with_next=False)
    signature.paragraph_format.keep_together = True


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style.font.size = Pt(20)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_borders(title_style)

    heading_style = doc.styles["Heading 1"]
    heading_style.font.name = "Arial"
    heading_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    heading_style.font.size = Pt(13)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(4)
    title.add_run("Chat WhatsApp Individual Semua Speaker")
    remove_paragraph_borders(title)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    set_run_font(
        subtitle.add_run(
            "Pesan undangan yang telah disesuaikan dengan penerima, jadwal, dan topik setiap sesi "
            "Founder's Live CEO Masterclass 2026."
        ),
        size=10.5,
        color="404040",
    )

    for index, speaker in enumerate(SPEAKERS, start=1):
        add_message(doc, index, speaker)
        if index % 2 == 0 and index < len(SPEAKERS):
            doc.add_page_break()

    doc.core_properties.title = "Chat WhatsApp Individual Semua Speaker"
    doc.core_properties.subject = "Undangan Founder's Live CEO Masterclass 2026"
    doc.core_properties.author = "Tim CEO Masterclass"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
