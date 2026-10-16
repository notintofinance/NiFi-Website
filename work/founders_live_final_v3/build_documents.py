from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path("/Users/agatbe/Claude Project/Claude Vibecode (NiFi)")
LETTER_SOURCE = ROOT / "work/updated_founders_live_v2/docx"
TOR_SOURCE = Path(
    "/Users/agatbe/Downloads/CMC Founder's Live - ToR & Undangan Speaker/"
    "Revisi 9 Okt 2026/CMC - ToR Founders Live 2026.docx"
)
POSTER = Path(
    "/var/folders/12/kf33qfk559sgsvh7429ht9bc0000gn/T/"
    "codex-clipboard-14fa46ce-a2ba-4262-9da6-e1aa4a5aa959.png"
)
OUT_ROOT = ROOT / "work/founders_live_final_v3"
LETTER_OUT = OUT_ROOT / "letters_docx"
TOR_OUT = OUT_ROOT / "tor_docx" / "CMC - ToR Founders Live 2026 - Umum.docx"


def set_paragraph_text_keep_format(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell_text_keep_format(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text_keep_format(paragraph, text)
    for extra in cell.paragraphs[1:]:
        set_paragraph_text_keep_format(extra, "")


def remove_all_highlights(doc: Document) -> None:
    for element in doc.element.xpath(".//w:highlight"):
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def drawing_paragraphs(doc: Document):
    return [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]


def replace_profile_visuals(doc: Document) -> None:
    pictures = drawing_paragraphs(doc)
    if len(pictures) < 2:
        raise RuntimeError(f"Expected at least two inline document images, found {len(pictures)}")

    first = pictures[0]
    for run in list(first.runs):
        first._p.remove(run._r)
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.add_run().add_picture(str(POSTER), width=Inches(4.30))

    second = pictures[1]
    second_parent = second._p.getparent()
    second_parent.remove(second._p)


def add_kadin_statement(doc: Document) -> None:
    old = (
        "CEO Masterclass adalah program pembelajaran bisnis yang materinya dibawakan oleh para "
        "pemimpin usaha nasional, antara lain Anindya Bakrie, Shinta Kamdani, Sandiaga Uno, dan "
        "Pandu Sjahrir, dengan dukungan Kadin Indonesia."
    )
    new = (
        "CEO Masterclass adalah bagian dari Kamar Dagang dan Industri (KADIN) Indonesia dan "
        "merupakan program pembelajaran bisnis yang materinya dibawakan oleh para pemimpin usaha "
        "nasional, antara lain Anindya Bakrie, Shinta Kamdani, Sandiaga Uno, dan Pandu Sjahrir."
    )
    for paragraph in doc.paragraphs:
        if old in paragraph.text:
            set_paragraph_text_keep_format(paragraph, paragraph.text.replace(old, new))
            return
    raise RuntimeError("Target CEO Masterclass paragraph not found")


def update_letters() -> list[Path]:
    LETTER_OUT.mkdir(parents=True, exist_ok=True)
    outputs = []
    for source in sorted(LETTER_SOURCE.glob("*.docx")):
        target = LETTER_OUT / source.name
        shutil.copy2(source, target)
        doc = Document(target)
        remove_all_highlights(doc)
        add_kadin_statement(doc)
        replace_profile_visuals(doc)
        doc.save(target)
        outputs.append(target)
    if len(outputs) != 9:
        raise RuntimeError(f"Expected 9 invitation documents, found {len(outputs)}")
    return outputs


def update_tor() -> Path:
    TOR_OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TOR_SOURCE, TOR_OUT)
    doc = Document(TOR_OUT)

    paragraph_replacements = {
        "8 sesi online · Selasa dan Kamis, 3–26 November 2026 · 19.00–20.30 WIB · Zoom":
            "8 sesi online · Selasa dan Kamis, 17 November-10 Desember 2026 · 19.00-20.30 WIB · Zoom",
    }
    for paragraph in doc.paragraphs:
        if paragraph.text in paragraph_replacements:
            set_paragraph_text_keep_format(paragraph, paragraph_replacements[paragraph.text])

    schedule_table = doc.tables[0]
    set_cell_text_keep_format(
        schedule_table.cell(0, 1),
        "8 sesi, setiap Selasa dan Kamis, 17 November-10 Desember 2026",
    )

    schedule_dates = [
        "Sel, 17 Nov",
        "Kam, 19 Nov",
        "Sel, 24 Nov",
        "Kam, 26 Nov",
        "Sel, 1 Des",
        "Kam, 3 Des",
        "Sel, 8 Des",
        "Kam, 10 Des",
    ]
    topic_table = doc.tables[2]
    for row_index, date_text in enumerate(schedule_dates, start=1):
        set_cell_text_keep_format(topic_table.cell(row_index, 0), date_text)

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Narasumber menerima honorarium sebesar"):
            set_paragraph_text_keep_format(
                paragraph,
                "Narasumber menerima honorarium sesuai kesepakatan tertulis dengan CEO "
                "Masterclass. Ketentuan nilai, pajak, dan waktu pembayaran disampaikan secara "
                "terpisah. Jika jadwal perlu berubah, kedua pihak akan mencari tanggal pengganti "
                "bersama.",
            )
            break
    else:
        raise RuntimeError("Honorarium paragraph not found")

    contact_table = doc.tables[4]
    set_cell_text_keep_format(
        contact_table.cell(0, 1),
        "Mario · +6285178906100 · ceomasterclassoffical@gmail.com",
    )
    set_cell_text_keep_format(
        contact_table.cell(1, 1),
        "ceomasterclassoffical@gmail.com · +6285178906100",
    )

    pictures = drawing_paragraphs(doc)
    if len(pictures) != 1:
        raise RuntimeError(f"Expected one body image in ToR, found {len(pictures)}")
    picture_paragraph = pictures[0]
    for run in list(picture_paragraph.runs):
        picture_paragraph._p.remove(run._r)
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(POSTER), width=Inches(4.35))

    remove_all_highlights(doc)
    doc.save(TOR_OUT)
    return TOR_OUT


def main() -> None:
    if not POSTER.exists():
        raise FileNotFoundError(POSTER)
    letters = update_letters()
    tor = update_tor()
    print(f"Updated {len(letters)} invitation DOCX files")
    print(f"Updated common ToR: {tor}")


if __name__ == "__main__":
    main()
