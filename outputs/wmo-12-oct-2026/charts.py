"""Native chart XML + embedded workbooks, styled like the template (Calibri 8pt, black axes, white chart area)."""
import datetime as dt
from xml.sax.saxutils import escape as esc
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L

NAVY, BLUE, ORANGE, GREY = "002060", "0051A2", "ED7D31", "A5A5A5"
NS = ('xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def serial(d):
    return (d - dt.date(1899, 12, 30)).days


def txpr(sz=800, b=0, color="000000", rot=None):
    bp = f'<a:bodyPr rot="{rot}" vert="horz"/>' if rot is not None else "<a:bodyPr/>"
    return (f'<c:txPr>{bp}<a:lstStyle/><a:p><a:pPr><a:defRPr sz="{sz}" b="{b}" i="0">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:latin typeface="Calibri"/></a:defRPr></a:pPr>'
            f'<a:endParaRPr lang="en-US"/></a:p></c:txPr>')


AX_LN = '<c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:prstDash val="solid"/></a:ln></c:spPr>'
NO_LN = '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'


def str_cache(ref, vals):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{esc(str(v))}</c:v></c:pt>' for i, v in enumerate(vals))
    return f'<c:strRef><c:f>{ref}</c:f><c:strCache><c:ptCount val="{len(vals)}"/>{pts}</c:strCache></c:strRef>'


def num_cache(ref, vals, fmt="General"):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals) if v is not None)
    return (f'<c:numRef><c:f>{ref}</c:f><c:numCache><c:formatCode>{esc(fmt)}</c:formatCode>'
            f'<c:ptCount val="{len(vals)}"/>{pts}</c:numCache></c:numRef>')


def lbl(idx, color, fmt, pos="t", dx=None, dy=None, sz=800):
    lay = (f'<c:layout><c:manualLayout><c:x val="{dx}"/><c:y val="{dy}"/></c:manualLayout></c:layout>'
           if dx is not None else "")
    return (f'<c:dLbl><c:idx val="{idx}"/>{lay}<c:numFmt formatCode="{esc(fmt)}" sourceLinked="0"/>'
            f'<c:spPr><a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>'
            f'{txpr(sz, 1, "FFFFFF")}<c:dLblPos val="{pos}"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
            f'<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>')


def dlbls(labels):
    return (f'<c:dLbls>{"".join(labels)}{NO_LN}<c:showLegendKey val="0"/><c:showVal val="0"/><c:showCatName val="0"/>'
            f'<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def line_ser(i, name, col, cats_xml, vals, color, width=19050, dash="solid", labels=(), fmt="General"):
    n = len(vals)
    return (f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/><c:tx>{str_cache(f"Sheet1!${col}$1", [name])}</c:tx>'
            f'<c:spPr><a:ln w="{width}" cap="rnd"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:prstDash val="{dash}"/><a:round/></a:ln></c:spPr><c:marker><c:symbol val="none"/></c:marker>'
            f'{dlbls(labels) if labels else ""}<c:cat>{cats_xml}</c:cat>'
            f'<c:val>{num_cache(f"Sheet1!${col}$2:${col}${n+1}", vals, fmt)}</c:val><c:smooth val="0"/></c:ser>')


def date_ax(axid, cross, lo, hi, fmt, unit, tunit, delete=0, pos="b"):
    return (f'<c:dateAx><c:axId val="{axid}"/><c:scaling><c:orientation val="minMax"/><c:max val="{hi}"/>'
            f'<c:min val="{lo}"/></c:scaling><c:delete val="{delete}"/><c:axPos val="{pos}"/>'
            f'<c:numFmt formatCode="{esc(fmt)}" sourceLinked="0"/><c:majorTickMark val="out"/><c:minorTickMark val="none"/>'
            f'<c:tickLblPos val="low"/>{AX_LN}{txpr()}<c:crossAx val="{cross}"/><c:crosses val="autoZero"/>'
            f'<c:auto val="1"/><c:lblOffset val="100"/><c:baseTimeUnit val="days"/><c:majorUnit val="{unit}"/>'
            f'<c:majorTimeUnit val="{tunit}"/></c:dateAx>')


def cat_ax(axid, cross, skip=1, pos="b", lblpos="low", orient="minMax", sz=800):
    return (f'<c:catAx><c:axId val="{axid}"/><c:scaling><c:orientation val="{orient}"/></c:scaling><c:delete val="0"/>'
            f'<c:axPos val="{pos}"/><c:numFmt formatCode="General" sourceLinked="0"/><c:majorTickMark val="none"/>'
            f'<c:minorTickMark val="none"/><c:tickLblPos val="{lblpos}"/>{AX_LN}{txpr(sz)}<c:crossAx val="{cross}"/>'
            f'<c:crosses val="autoZero"/><c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/>'
            f'<c:tickLblSkip val="{skip}"/><c:tickMarkSkip val="{skip}"/><c:noMultiLvlLbl val="0"/></c:catAx>')


def val_ax(axid, cross, lo, hi, unit, fmt, pos="l", delete=0, grid=False, crosses="autoZero"):
    g = ('<c:majorGridlines><c:spPr><a:ln w="6350"><a:solidFill><a:srgbClr val="D9D9D9"/></a:solidFill>'
         '</a:ln></c:spPr></c:majorGridlines>') if grid else ""
    return (f'<c:valAx><c:axId val="{axid}"/><c:scaling><c:orientation val="minMax"/><c:max val="{hi}"/>'
            f'<c:min val="{lo}"/></c:scaling><c:delete val="{delete}"/><c:axPos val="{pos}"/>{g}'
            f'<c:numFmt formatCode="{esc(fmt)}" sourceLinked="0"/><c:majorTickMark val="out"/><c:minorTickMark val="none"/>'
            f'<c:tickLblPos val="nextTo"/>{AX_LN}{txpr()}<c:crossAx val="{cross}"/><c:crosses val="{crosses}"/>'
            f'<c:crossBetween val="between"/><c:majorUnit val="{unit}"/></c:valAx>')


def legend(h=0.12):
    return (f'<c:legend><c:legendPos val="t"/><c:layout><c:manualLayout><c:xMode val="edge"/><c:yMode val="edge"/>'
            f'<c:x val="0"/><c:y val="0"/><c:w val="1"/><c:h val="{h}"/></c:manualLayout></c:layout>'
            f'<c:overlay val="0"/>{txpr(800, 1)}</c:legend>')


def space(plot, layout, leg=""):
    x, y, w, h = layout
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<c:chartSpace {NS}><c:date1904 val="0"/>'
            f'<c:lang val="en-US"/><c:roundedCorners val="0"/><c:style val="2"/><c:chart><c:autoTitleDeleted val="1"/>'
            f'<c:plotArea><c:layout><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
            f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/></c:manualLayout></c:layout>{plot}{NO_LN}'
            f'</c:plotArea>{leg}<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>'
            f'<c:spPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>'
            f'{txpr()}<c:externalData r:id="rId1"><c:autoUpdate val="0"/></c:externalData></c:chartSpace>')


def workbook(path, header, rows, date_col=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    if date_col:
        for c in ws["A"][1:]:
            c.number_format = "dd-mmm-yy"
    ws.column_dimensions["A"].width = 14
    wb.save(path)


# ------------------------------------------------------------------ chart types
def line_chart(dates, series, ylo, yhi, yunit, yfmt, xfmt, xunit, xtunit, layout, show_legend=True, leg_h=0.12):
    """series: list of (name, values, color, width, dash, labels[(idx, dx, dy)], lblfmt)."""
    n = len(dates)
    cats = num_cache(f"Sheet1!$A$2:$A${n+1}", [serial(d) for d in dates], "dd-mmm-yy")
    sers = ""
    for i, (name, vals, color, width, dash, labels, lf) in enumerate(series):
        col = L(i + 2)
        lab = [lbl(ix, color, lf, "t", dx, dy) for ix, dx, dy in labels]
        sers += line_ser(i, name, col, cats, vals, color, width, dash, lab)
    plot = (f'<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>{sers}<c:marker val="1"/>'
            f'<c:axId val="111"/><c:axId val="222"/></c:lineChart>'
            + date_ax(111, 222, serial(dates[0]), serial(dates[-1]), xfmt, xunit, xtunit)
            + val_ax(222, 111, ylo, yhi, yunit, yfmt))
    return space(plot, layout, legend(leg_h) if show_legend else "")


def cat_line_chart(cats, series, ylo, yhi, yunit, yfmt, skip, layout, leg_h=0.12):
    n = len(cats)
    cx = str_cache(f"Sheet1!$A$2:$A${n+1}", cats)
    sers = ""
    for i, (name, vals, color, width, dash, labels, lf) in enumerate(series):
        col = L(i + 2)
        lab = [lbl(ix, color, lf, "t", dx, dy) for ix, dx, dy in labels]
        sers += line_ser(i, name, col, cx, vals, color, width, dash, lab)
    plot = (f'<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>{sers}<c:marker val="1"/>'
            f'<c:axId val="111"/><c:axId val="222"/></c:lineChart>'
            + cat_ax(111, 222, skip) + val_ax(222, 111, ylo, yhi, yunit, yfmt))
    return space(plot, layout, legend(leg_h))


def hbar_chart(cats, vals, colors, xlo, xhi, xunit, fmt, layout):
    n = len(cats)
    dpts = "".join(f'<c:dPt><c:idx val="{i}"/><c:invertIfNegative val="0"/><c:bubble3D val="0"/>'
                   f'<c:spPr><a:solidFill><a:srgbClr val="{c}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr></c:dPt>'
                   for i, c in enumerate(colors))
    lbls = ('<c:dLbls>' + f'<c:numFmt formatCode="{esc(fmt)}" sourceLinked="0"/>' + NO_LN + txpr(800, 1)
            + '<c:dLblPos val="outEnd"/><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/>'
            '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')
    ser = (f'<c:ser><c:idx val="0"/><c:order val="0"/><c:tx>{str_cache("Sheet1!$B$1", ["WoW return (%)"])}</c:tx>'
           f'<c:spPr><a:solidFill><a:srgbClr val="{BLUE}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>'
           f'<c:invertIfNegative val="0"/>{dpts}{lbls}<c:cat>{str_cache(f"Sheet1!$A$2:$A${n+1}", cats)}</c:cat>'
           f'<c:val>{num_cache(f"Sheet1!$B$2:$B${n+1}", vals, "0.00")}</c:val></c:ser>')
    plot = (f'<c:barChart><c:barDir val="bar"/><c:grouping val="clustered"/><c:varyColors val="0"/>{ser}'
            f'<c:gapWidth val="45"/><c:axId val="333"/><c:axId val="444"/></c:barChart>'
            + cat_ax(333, 444, 1, pos="l", lblpos="low", orient="maxMin")
            + val_ax(444, 333, xlo, xhi, xunit, "0", pos="b", delete=1))
    return space(plot, layout)


def bar_line_combo(cats, bar_name, bar_vals, line_name, line_vals, left, right, layout, bar_fmt="0.0", line_fmt="0.0",
                   line_labels=(), bar_color=None, line_color=None):
    """Clustered columns on the left axis + line on the right axis (template colours)."""
    bar_color = bar_color or BLUE
    line_color = line_color or ORANGE
    n = len(cats)
    cx = str_cache(f"Sheet1!$A$2:$A${n+1}", cats)
    blbl = ('<c:dLbls>' + f'<c:numFmt formatCode="{esc(bar_fmt)}" sourceLinked="0"/>' + NO_LN + txpr(700, 0)
            + '<c:dLblPos val="outEnd"/><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/>'
            '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')
    bser = (f'<c:ser><c:idx val="0"/><c:order val="0"/><c:tx>{str_cache("Sheet1!$B$1", [bar_name])}</c:tx>'
            f'<c:spPr><a:solidFill><a:srgbClr val="{bar_color}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>'
            f'<c:invertIfNegative val="0"/>{blbl}<c:cat>{cx}</c:cat>'
            f'<c:val>{num_cache(f"Sheet1!$B$2:$B${n+1}", bar_vals, "0.00")}</c:val></c:ser>')
    lab = [lbl(ix, line_color, line_fmt, "t", dx, dy) for ix, dx, dy in line_labels]
    lser = line_ser(1, line_name, "C", cx, line_vals, line_color, 22225, "solid", lab)
    plot = (f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:varyColors val="0"/>{bser}'
            f'<c:gapWidth val="60"/><c:axId val="11"/><c:axId val="22"/></c:barChart>'
            f'<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>{lser}<c:marker val="1"/>'
            f'<c:axId val="33"/><c:axId val="44"/></c:lineChart>'
            + cat_ax(11, 22, 1)
            + val_ax(22, 11, left[0], left[1], left[2], left[3])
            + cat_ax(33, 44, 1).replace('<c:delete val="0"/>', '<c:delete val="1"/>')
            + val_ax(44, 33, right[0], right[1], right[2], right[3], pos="r", crosses="max"))
    return space(plot, layout, legend(0.14))


def dual_line_chart(dates, left, right, laxis, raxis, xfmt, xunit, xtunit, layout, leg_h=0.12):
    """Two lines on one date axis: `left` on the left value axis, `right` on the right value axis.
    left/right: (name, values, color, labels[(idx, dx, dy)], lblfmt); axis: (lo, hi, unit, fmt)."""
    n = len(dates)
    cats = num_cache(f"Sheet1!$A$2:$A${n+1}", [serial(d) for d in dates], "dd-mmm-yy")
    ln, lv, lc, ll, lf = left
    rn, rv, rc, rl, rf = right
    s1 = line_ser(0, ln, "B", cats, lv, lc, 22225, "solid", [lbl(i, lc, lf, "t", dx, dy) for i, dx, dy in ll])
    s2 = line_ser(1, rn, "C", cats, rv, rc, 22225, "solid", [lbl(i, rc, rf, "t", dx, dy) for i, dx, dy in rl])
    lo, hi = serial(dates[0]), serial(dates[-1])
    plot = (f'<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>{s1}<c:marker val="1"/>'
            f'<c:axId val="111"/><c:axId val="222"/></c:lineChart>'
            f'<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>{s2}<c:marker val="1"/>'
            f'<c:axId val="333"/><c:axId val="444"/></c:lineChart>'
            + date_ax(111, 222, lo, hi, xfmt, xunit, xtunit)
            + val_ax(222, 111, laxis[0], laxis[1], laxis[2], laxis[3])
            + date_ax(333, 444, lo, hi, xfmt, xunit, xtunit, delete=1)
            + val_ax(444, 333, raxis[0], raxis[1], raxis[2], raxis[3], pos="r", crosses="max"))
    return space(plot, layout, legend(leg_h))
