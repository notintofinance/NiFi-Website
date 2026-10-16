import openpyxl, datetime as dt, re
F="/Users/agatbe/Downloads/TIM - 9-3.xlsx"
def clean(h): return re.sub(r'\s*\((L|R)\d\)\s*$','',str(h)).strip()
def load():
    wb=openpyxl.load_workbook(F,data_only=True); out={}
    for ws in wb:
        hdr={c:ws.cell(1,c).value for c in range(1,ws.max_column+1)}
        c=1
        while c<=ws.max_column:
            if hdr.get(c)=='Date':
                vc=c+1
                while vc<=ws.max_column and hdr.get(vc) not in (None,'Date'):
                    ser={}
                    for r in range(2,ws.max_row+1):
                        d=ws.cell(r,c).value; v=ws.cell(r,vc).value
                        if isinstance(d,dt.datetime) and isinstance(v,(int,float)): ser[d.date()]=float(v)
                    out[f"{ws.title}|{clean(hdr[vc])}|dcol{c}"]=ser; vc+=1
                c=vc
            else: c+=1
    return out
