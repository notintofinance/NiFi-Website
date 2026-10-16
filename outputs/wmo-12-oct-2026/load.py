"""Parse TIM-8: each 'Date' column owns the value columns to its right until a blank/next Date."""
import openpyxl, datetime as dt, re
F="/Users/agatbe/Downloads/TIM - 8.xlsx"
def clean(h):
    return re.sub(r'\s*\((L|R)\d\)\s*$','',str(h)).strip()
def load():
    wb=openpyxl.load_workbook(F,data_only=True)
    out={}
    for ws in wb:
        hdr={c:ws.cell(2,c).value for c in range(1,ws.max_column+1)}
        c=1
        while c<=ws.max_column:
            if hdr.get(c)=='Date':
                vc=c+1
                while vc<=ws.max_column and hdr.get(vc) not in (None,'Date'):
                    name=f"{ws.title}|{clean(hdr[vc])}|dcol{c}"
                    ser={}
                    for r in range(3,ws.max_row+1):
                        d=ws.cell(r,c).value; v=ws.cell(r,vc).value
                        if isinstance(d,dt.datetime) and isinstance(v,(int,float)):
                            ser[d.date()]=float(v)
                    out[name]=ser
                    vc+=1
                c=vc
            else: c+=1
    return out
if __name__=='__main__':
    S=load()
    for k,v in S.items():
        ds=sorted(v); gaps=sorted({(b-a).days for a,b in zip(ds,ds[1:])})
        print(f"{k:60s} n={len(v):4d} {ds[0]}..{ds[-1]} gaps={gaps[:5]} last={v[ds[-1]]}")
