from pathlib import Path

from pptx import Presentation
from pptx.util import Emu
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
PPTX_IMAGES = sorted((ROOT / "highres_from_pptx").glob("slide-*.png"))
PDF_IMAGES = sorted((ROOT / "highres_from_pdf").glob("page-*.png"))

# Exact portrait slide size from the source PPTX.
SLIDE_W = 7_556_500
SLIDE_H = 10_693_400


def build_pptx(images: list[Path], output: Path) -> None:
    deck = Presentation()
    deck.slide_width = Emu(SLIDE_W)
    deck.slide_height = Emu(SLIDE_H)
    blank = deck.slide_layouts[6]
    for image in images:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image),
            0,
            0,
            width=deck.slide_width,
            height=deck.slide_height,
        )
    deck.core_properties.title = "Partnership NIFI x ICMSS — Export-Safe"
    deck.core_properties.subject = "Flattened high-resolution version to prevent stitching and clipping artifacts"
    deck.save(output)


def build_pdf(images: list[Path], output: Path) -> None:
    # Preserve the source A4 aspect ratio while using the source PPTX's exact dimensions.
    page_w = 595.0
    page_h = page_w * SLIDE_H / SLIDE_W
    pdf = canvas.Canvas(str(output), pagesize=(page_w, page_h), pageCompression=1)
    pdf.setTitle(output.stem)
    for image in images:
        pdf.drawImage(str(image), 0, 0, width=page_w, height=page_h, preserveAspectRatio=False)
        pdf.showPage()
    pdf.save()


OUT.mkdir(exist_ok=True)
build_pptx(PPTX_IMAGES, OUT / "Partnership NIFI x ICMSS - Export-Safe.pptx")
build_pdf(PPTX_IMAGES, OUT / "Partnership NIFI x ICMSS - Export-Safe.pdf")
build_pdf(PDF_IMAGES, OUT / "Partnership NIFI x ICMSS-2 - Fixed.pdf")
