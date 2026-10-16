from pathlib import Path
import zipfile

from pptx import Presentation
from pptx.util import Emu


ROOT = Path(__file__).resolve().parent
IMAGES = sorted((ROOT / "highres_from_pptx").glob("slide-*.png"))
OUT = ROOT / "split_decks"
SLIDE_W = 7_556_500
SLIDE_H = 10_693_400

# Each focused deck retains the shared company/program introduction (slides 1–12)
# and the common closing slide (25).
DECKS = {
    # Full master deck containing every partnership route except slide 23.
    "NIFI x ICMSS - MASTER All Partnership Packages.pptx": list(range(1, 23)) + [24, 25],
    "NIFI x ICMSS - Talent Partnership.pptx": list(range(1, 16)) + [25],
    "NIFI x ICMSS - Case Collaborator.pptx": list(range(1, 13)) + [16, 17, 18, 25],
    # Slide 23 (Sponsorship Package tier table) is intentionally omitted.
    "NIFI x ICMSS - Media and Community Partnership.pptx": list(range(1, 13)) + [19, 20, 21, 22, 24, 25],
}


def build_deck(name: str, slide_numbers: list[int]) -> Path:
    deck = Presentation()
    deck.slide_width = Emu(SLIDE_W)
    deck.slide_height = Emu(SLIDE_H)
    blank = deck.slide_layouts[6]
    for number in slide_numbers:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(IMAGES[number - 1]),
            0,
            0,
            width=deck.slide_width,
            height=deck.slide_height,
        )
    deck.core_properties.title = Path(name).stem
    deck.core_properties.subject = "Focused export-safe partnership proposal"
    output = OUT / name
    deck.save(output)
    return output


OUT.mkdir(exist_ok=True)
outputs = [build_deck(name, slides) for name, slides in DECKS.items()]
zip_path = ROOT / "NIFI x ICMSS - Focused Partnership Decks.zip"
with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
    for output in outputs:
        archive.write(output, arcname=output.name)

print(zip_path)
for output in outputs:
    print(output)
